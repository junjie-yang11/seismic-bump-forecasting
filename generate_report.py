"""Generate the report from result CSVs, with one source for Markdown and Word.

Run experiments first. Word creation needs requirements-report.txt. On Windows,
export_report.ps1 exports the Word report to PDF using installed Microsoft Word.
"""
from __future__ import annotations
import csv
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent

def rows(name):
    with (ROOT / 'results' / name).open(encoding='utf-8', newline='') as f:
        return list(csv.DictReader(f))

def short(model):
    return 'LR' if 'logistic' in model else 'CART'

def table(headers, data):
    return '\n'.join(['| ' + ' | '.join(headers) + ' |',
        '| ' + ' | '.join(['---'] * len(headers)) + ' |'] +
        ['| ' + ' | '.join(map(str, row)) + ' |' for row in data])

def number(value):
    return '{:.4f}'.format(float(value))

def report_markdown():
    met, hold = rows('metrics_by_scheme.csv'), rows('holdout.csv')
    shared, cal = rows('shared_test_comparison.csv'), rows('calibration.csv')
    integrity = json.loads((ROOT / 'results/integrity.json').read_text(encoding='utf-8'))
    parts = [
        '# Validation and calibration of seismic bump forecasting models',
        'Junjie Yang\n\nMining Engineering, Fuzhou University',
        'Revised 2 October 2026. Code and result files: https://github.com/junjie-yang11/seismic-bump-forecasting',
        '## Abstract',
        'This study compares random stratified validation, validation in record order, '
        'and a chronological holdout on the UCI Seismic Bumps mirror. Performance varies '
        'with the protocol. The original comparison uses different test populations, '
        'so its PR-AUC gap combines test composition, prevalence drift, training history '
        'and possible temporal leakage. The same-test comparison below controls the test '
        'rows but does not isolate a causal leakage effect. Calibration priors and alert '
        'thresholds are now derived separately inside each training fold.',
        '## 1 Data and scope',
        'The mirror contains {n_rows} records and {n_positive} positives. Overall prevalence '
        'is {prevalence:.4f}. There are no duplicate rows in this mirror. The task is '
        'prediction of hazardous seismic events in the next shift. Row order is used '
        'as a temporal proxy; explicit timestamps and longwall identifiers are unavailable. '
        'The results are conditional on this order representing a meaningful forecasting '
        'sequence. The mirror has six fewer rows than the original 2584-record dataset.'.format(**integrity),
        'The model uses 17 encoded features after removing nbumps. Three retained columns '
        'are constant: nbumps6, nbumps7 and nbumps89. Dropping nbumps does not make the '
        'design matrix full rank; L2 regularization and constant-safe standardization '
        'allow the logistic model to fit.',
        table(['Record block', '1', '2', '3', '4', '5'],
              [['Prevalence'] + [number(integrity['prevalence_by_quintile'][str(i)]) for i in range(5)]]),
        '![Figure 1 Hazard prevalence across record blocks](../results/figures/prevalence_drift.png)',
        '## 2 Models and validation',
        'Logistic regression is fitted with IRLS, L2 penalty 1.0, an unpenalized '
        'intercept and balanced class weights. Its feature standardization is fitted '
        'inside training only. CART uses 60 ordinary bootstrap trees, depth 6, minimum '
        '20 observations per leaf and unweighted Gini impurity. Its leaf probabilities '
        'are unweighted sample proportions; CART receives no balanced-prior correction.',
        'Random stratified validation uses ten folds and seeds 0 to 4. All metrics in '
        'its aggregate table are averaged over those five seeds. Record-order validation '
        'divides the data into five consecutive blocks and evaluates four expanding '
        'training windows. The first block is training only. Holdout uses the first '
        '70 percent for training and the last 30 percent for testing. Test blocks with '
        'zero positives are retained; single-class ROC-AUC is undefined rather than '
        'being used to suppress their predictions.',
        '## 3 Thresholds and calibration',
        'Each outer training fold has an inner split. Record-order folds reserve the '
        'last 20 percent of available training history for threshold reference; random '
        'folds reserve 20 percent of each training class with seed 7. An inner model '
        'predicts these reference rows. The alert budget is the positive proportion '
        'in its inner fit data. The threshold excludes the boundary score and all '
        'ties at that boundary, so the reference alert count never exceeds the integer '
        'budget. The outer model is refitted on its full training fold and evaluated '
        'with this fixed numerical threshold. Test labels and test score quantiles '
        'never select the threshold. Refitting and distribution drift can change the '
        'future alert rate; the budget is not a promise of an exact test alert rate.',
        'For weighted logistic regression, prior correction shifts log odds from '
        'effective prior 0.5 to that outer fold\'s training prevalence. Correction is '
        'performed before pooling the calibrated predictions. Its adequacy under '
        'distribution drift remains empirical. Ranking metrics below use raw scores. '
        'Brier skill uses the evaluation prevalence as a retrospective oracle constant '
        'reference, not as a deployable forecast or calibration prior.',
        '## 4 Validation results',
        table(['Model', 'Validation', 'ROC AUC', 'PR AUC', 'Recall', 'Alert rate'],
            [[short(r['model']), 'Random' if 'random' in r['validation'] else 'Record order'] +
             [number(r[k]) for k in ('roc_auc', 'pr_auc', 'recall', 'alert_rate')] for r in met]),
        'Accuracy, precision, F1 and balanced accuracy are included in '
        'results/metrics_by_scheme.csv; seed-level metrics are in random_metrics_by_seed.csv.',
        '![Figure 2 PR AUC by model and validation scheme](../results/figures/pr_auc_by_scheme.png)',
        'The two schemes evaluate different populations. Random validation covers all '
        'records, whereas record-order validation excludes the first training-only '
        'block. Their positive proportions differ. The original approximately 57 percent '
        'LR PR-AUC reduction cannot be assigned wholly to temporal leakage.',
        '### Same test rows',
        table(['Model', 'Test n', 'Prevalence', 'Random PR AUC', 'Order PR AUC'],
              [[short(r['model']), str(int(float(r['n']))), number(r['prevalence']),
                number(r['random_pr_auc']), number(r['time_pr_auc'])] for r in shared]),
        'Random scores are restricted to the exact rows covered by record-order '
        'validation, and averaged over the same five seeds. This controls test '
        'composition. Training sizes, permitted training periods and cross-fold '
        'score scales still differ; this is a sensitivity analysis, not an estimate '
        'of a pure leakage percentage. Per-seed results are in shared_test_by_seed.csv.',
        '![Figure 3 Precision recall curves from seed 0 and record order](../results/figures/pr_curves.png)',
        '![Figure 4 ROC curves from seed 0 and record order](../results/figures/roc_curves.png)',
        '## 5 Holdout and probability calibration',
        table(['Model', 'ROC AUC', 'PR AUC', 'Recall', 'Precision', 'Alert rate'],
            [[short(r['model'])] + [number(r[k]) for k in
              ('roc_auc', 'pr_auc', 'recall', 'precision', 'alert_rate')] for r in hold]),
        'These holdout operating metrics use fixed thresholds selected on training '
        'validation rows. They replace the earlier retrospective test-quantile metrics.',
        table(['Model', 'Variant', 'Brier', 'ECE', 'Mean prediction', 'Observed'],
            [[r['model'], 'Fold prior correction' if 'correction' in r['variant'] else 'As fitted'] +
             [number(r[k]) for k in ('brier', 'ece', 'mean_predicted', 'observed_rate')] for r in cal]),
        'Calibration numbers in this table come directly from results/calibration.csv. '
        'MCE and Brier skill are available in that file. Only LR has a corrected variant.',
        '![Figure 5 Logistic regression calibration with fold local priors](../results/figures/reliability_lr.png)',
        '## 6 Record index probe and feature importance',
        table(['Model', 'Validation', 'ROC AUC', 'PR AUC'],
            [[short(r['model']), 'Random seed 0' if 'random' in r['validation'] else 'Record order',
              number(r['roc_auc']), number(r['pr_auc'])] for r in rows('leakage_probe.csv')]),
        'The record index is a proxy for sequence position, not a physical monitoring '
        'variable. Its effect suggests sensitivity to record structure but does not '
        'prove a unique leakage mechanism. Both random probe and corresponding raw '
        'score reference should be compared at seed 0.',
        table(['Feature', 'In sample ROC AUC drop', 'Standard deviation'],
            [[r['feature'], number(r['auc_drop']), number(r['sd'])] for r in rows('importance.csv')[:8]]),
        'Permutation importance is measured on the data used to fit the forest and '
        'is descriptive. It is not independent evidence of future feature usefulness. '
        'Correlated monitoring features can mask one another\'s contributions.',
        '## 7 Limitations and reproducibility',
        'No timestamp or wall identifier confirms the ordering assumption. No '
        'hyperparameter search or confidence interval for temporally correlated '
        'observations is provided. Inner model thresholds are transferred to a '
        'refitted outer model, so score-scale changes remain a limitation. Prior '
        'correction alone does not establish reliable deployment probabilities.',
        'Run python run_experiments.py, python -m unittest discover -s tests -v, '
        'and python verify_results.py. The verifier recomputes metrics from stored '
        'predictions, checks training-only priors and fixed thresholds, tests required '
        'row coverage and compares generated report text with result files. It does '
        'not claim that numerical agreement proves the scientific assumptions.',
        'Run python generate_report.py --docx with requirements-report.txt installed, '
        'then run export_report.ps1 on Windows with Microsoft Word installed. Markdown '
        'and Word are generated from the same report source; the PDF is exported from Word.',
        '## References',
        'Sikora M and Wrobel L (2010). Application of rule induction algorithms for '
        'analysis of data collected by seismic hazard monitoring systems in coal mines. '
        'Archives of Mining Sciences 55(1), 91-114.',
        'UCI Machine Learning Repository. Seismic Bumps dataset. '
        'https://archive.ics.uci.edu/ml/datasets/seismic-bumps',
        'Bergmeir C and Benitez JM (2012). On the use of cross-validation for time series '
        'predictor evaluation. Information Sciences 191, 192-213.',
        'Roberts DR et al (2017). Cross-validation strategies for data with temporal, '
        'spatial, hierarchical, or phylogenetic structure. Ecography 40, 913-929.',
    ]
    return '\n\n'.join(parts) + '\n'

def write_markdown_report():
    (ROOT / 'report/technical_note.md').write_text(report_markdown(), encoding='utf-8')

def write_word_report():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption'):
        s = doc.styles[name]
        s.font.name = 'Times New Roman'
        s.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Heading 1'].font.size = Pt(15)
    doc.styles['Heading 2'].font.size = Pt(12)
    doc.styles['Caption'].font.size = Pt(9)
    source = report_markdown()
    lines = source.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line:
            i += 1; continue
        if line.startswith('|'):
            data = []
            while i < len(lines) and lines[i].startswith('|'):
                data.append([v.strip() for v in lines[i].strip('|').split('|')]); i += 1
            data.pop(1)
            tbl = doc.add_table(rows=0, cols=len(data[0]))
            tbl.autofit = False
            width = 6.8 / len(data[0])
            for col in tbl.columns: col.width = Inches(width)
            for ri, values in enumerate(data):
                cells = tbl.add_row().cells
                for cell, value in zip(cells, values):
                    cell.text = value
                    cell.width = Inches(width)
                    pr = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement('w:tcBorders')
                    for edge in ('top', 'left', 'bottom', 'right'):
                        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single')
                        e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
                    pr.append(borders)
                    if ri == 0:
                        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'E8EDF2'); pr.append(shading)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(4)
                        p.paragraph_format.space_after = Pt(4)
                        for r in p.runs: r.font.size = Pt(9); r.bold = ri == 0
                trpr = tbl.rows[-1]._tr.get_or_add_trPr()
                trpr.append(OxmlElement('w:cantSplit'))
                if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        image = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
        if image:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str((ROOT / 'report' / image.group(2)).resolve()), width=Inches(5.3))
            caption = doc.add_paragraph(image.group(1), style='Caption')
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('# '): doc.add_paragraph(line[2:], style='Title')
        elif line.startswith('## '): doc.add_paragraph(line[3:], style='Heading 1')
        elif line.startswith('### '): doc.add_paragraph(line[4:], style='Heading 2')
        else: doc.add_paragraph(line)
        i += 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Seismic bump forecasting | ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    doc.save(str(ROOT / 'report/technical_report.docx'))

if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser(); ap.add_argument('--docx', action='store_true')
    args = ap.parse_args()
    write_markdown_report()
    if args.docx: write_word_report()
