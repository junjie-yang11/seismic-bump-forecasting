"""Independent saved-versus-fresh comparison; never edits experiment evidence."""
import argparse, ast, hashlib, json, re, subprocess, sys
from datetime import datetime, timezone
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document
from pypdf import PdfReader

def digest(path): return hashlib.sha256(path.read_bytes()).hexdigest()

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root',type=Path,required=True)
    parser.add_argument('--baseline',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args(); root=args.root.resolve(); baseline=args.baseline.resolve()
    record={'scope':'Numerical replication, report agreement, source parsing and current navigation; no claim of field validity',
            'created_at_utc':datetime.now(timezone.utc).isoformat(),'rtol':1e-9,'atol':1e-10,
            'csv':[],'npz':[],'native_models':[],'reports':[],'source_inventory':[],'links_checked':0,'passed':False}
    try:
        for path in sorted((root/'results').rglob('*.csv')):
            old=baseline/path.relative_to(root)
            if not old.exists(): raise AssertionError('Unexpected new numeric result '+str(path.relative_to(root)))
            left,right=pd.read_csv(old),pd.read_csv(path)
            if left.shape!=right.shape or list(left.columns)!=list(right.columns): raise AssertionError('CSV shape/schema '+path.name)
            maximum=0.; differing=0
            for column in left:
                if pd.api.types.is_numeric_dtype(left[column]) and pd.api.types.is_numeric_dtype(right[column]):
                    a,b=left[column].to_numpy(),right[column].to_numpy()
                    if not np.allclose(a,b,rtol=1e-9,atol=1e-10,equal_nan=True): raise AssertionError('CSV numeric '+path.name+' '+column)
                    finite=np.isfinite(a)&np.isfinite(b)
                    if finite.any(): maximum=max(maximum,float(np.max(np.abs(a[finite].astype(float)-b[finite].astype(float)))))
                    differing+=int(np.count_nonzero(~((a==b)|(pd.isna(a)&pd.isna(b)))))
                    if re.search(r'(^row$|row_|_row|^fold$|^seed$|^alerts$|^tp$|^fp$|^fn$|^tn$|^positives$|^excess$)',column):
                        if not np.array_equal(a,b,equal_nan=True): raise AssertionError('Discrete evidence '+path.name+' '+column)
                elif not left[column].equals(right[column]): raise AssertionError('CSV identities '+path.name+' '+column)
            record['csv'].append({'file':path.relative_to(root).as_posix(),'rows':len(right),'columns':len(right.columns),
                'byte_identical':digest(path)==digest(old),'max_absolute_numeric_difference':maximum,'differing_numeric_cells':differing})
        expected={p.relative_to(baseline).as_posix() for p in (baseline/'results').rglob('*.csv')}
        actual={p['file'] for p in record['csv']}
        if expected!=actual: raise AssertionError('Missing published CSV '+str(expected-actual))
        for path in (root/'results/phase1/research_models').glob('*.json'):
            if digest(path)!=digest(baseline/path.relative_to(root)): raise AssertionError('Native model '+path.name)
            record['native_models'].append({'file':path.relative_to(root).as_posix(),'byte_identical':True})
        for path in (root/'results').rglob('*.npz'):
            with np.load(path) as current,np.load(baseline/path.relative_to(root)) as old:
                if current.files!=old.files: raise AssertionError('NPZ names '+path.name)
                for key in current.files:
                    if not np.array_equal(current[key],old[key],equal_nan=True): raise AssertionError('NPZ replicates '+path.name+' '+key)
                record['npz'].append({'file':path.relative_to(root).as_posix(),'arrays':len(current.files),'all_arrays_exact':True})
        source=json.loads((root/'results/phase1/source_audit.json').read_text())
        previous=json.loads((baseline/'results/phase1/source_audit.json').read_text())
        if source!=previous: raise AssertionError('Source audit changed')
        record['source_audit']=source
        if digest(root/'phase2/locked_plan.json')!=digest(baseline/'phase2/locked_plan.json'): raise AssertionError('Locked plan changed')
        record['locked_plan_unchanged']=True
        before=json.loads((baseline/'results/phase2/run_manifest.json').read_text())['first_stage_hashes']
        after=json.loads((root/'results/phase2/run_manifest.json').read_text())['first_stage_hashes']
        record['first_stage_snapshot_keys']={'old_count':len(before),'fresh_count':len(after),
            'omitted':sorted(set(before)-set(after)),'added':sorted(set(after)-set(before))}
        for name,value in after.items():
            if digest(root/name)!=value: raise AssertionError('Stage 1 changed after Stage 2 '+name)
        for phase,name in [('phase1','technical_report'),('phase2','phase2_threshold_transfer_report')]:
            folder=root/'reports'/phase; doc=Document(folder/(name+'.docx')); pdf=PdfReader(folder/(name+'.pdf'))
            mdpath=folder/('technical_note.md' if phase=='phase1' else name+'.md')
            oldmd=baseline/mdpath.relative_to(root)
            unchanged=mdpath.read_bytes()==oldmd.read_bytes()
            if not unchanged: raise AssertionError('Unexpected manuscript change '+phase)
            old=Document(baseline/(folder/(name+'.docx')).relative_to(root))
            tables=lambda document:[[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
            if tables(doc)!=tables(old): raise AssertionError('Changed manuscript tables '+phase)
            text='\n'.join(page.extract_text() for page in pdf.pages)
            def numbers(value): return set(re.findall(r'-?\d+\.\d+',value.replace('−','-')))
            vals=numbers('\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells))
            if not vals.issubset(numbers(text)): raise AssertionError('PDF table numbers '+phase)
            if any(len(page.extract_text())<150 for page in pdf.pages): raise AssertionError('Blank PDF page '+phase)
            if doc.core_properties.author or doc.core_properties.last_modified_by: raise AssertionError('Private Word metadata '+phase)
            if '/Author' in (pdf.metadata or {}): raise AssertionError('Private PDF metadata '+phase)
            record['reports'].append({'stage':phase,'pages':len(pdf.pages),'tables':len(doc.tables),'figures':len(doc.inline_shapes),
                'manuscript_unchanged':unchanged,'tables_exact':True,'pdf_table_decimals_present':len(vals),
                'sha256':{p.name:digest(p) for p in [mdpath,folder/(name+'.docx'),folder/(name+'.pdf')]}})
        for path in sorted([*root.glob('src/*.py'),*root.glob('phase1/*.py'),*root.glob('phase2/*.py'),*root.glob('tests/*.py')]):
            ast.parse(path.read_text(encoding='utf-8'))
            record['source_inventory'].append({'file':path.relative_to(root).as_posix(),'sha256':digest(path)})
        guides=[root/'README.md',root/'CONTRIBUTIONS.md',*root.glob('docs/*.md'),*root.glob('*/README.md'),
                *root.glob('reports/**/*.md'),*root.glob('results/**/README.md')]
        auditguide=root/'docs/audits/2026-10-07-full-review/REVIEW.md'
        if auditguide.exists(): guides.append(auditguide)
        for guide in set(guides):
            for target in re.findall(r'\]\(([^)]+)\)',guide.read_text(encoding='utf-8')):
                if '://' in target or target.startswith(('mailto:','#')): continue
                name=target.split('#',1)[0]
                if name and not (guide.parent/name).exists(): raise AssertionError('Broken current link '+str(guide.relative_to(root))+' '+target)
                record['links_checked']+=1
        record['passed']=True
    except Exception as error:
        record['error_type']=type(error).__name__;record['error']=str(error)
        raise
    finally:
        args.output.parent.mkdir(parents=True,exist_ok=True)
        args.output.write_text(json.dumps(record,indent=2),encoding='utf-8')
    print('Release comparison passed:',len(record['csv']),'CSV files;',len(record['npz']),'NPZ files;',len(record['source_inventory']),'Python sources;',record['links_checked'],'current links')
    print('Snapshot key differences:',record['first_stage_snapshot_keys'])

if __name__=='__main__': main()
