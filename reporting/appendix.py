"""Generate a complete, printable supplement from frozen comparison outputs."""
import csv
from datetime import datetime,timezone
import hashlib
import json
from pathlib import Path
from reporting.overall import write_outputs
from reporting.edition import EDITION_TAG,edition_url

ROOT=Path(__file__).resolve().parents[1]
STEM='decision_comparison_supplement'
TITLE='Frozen Warning Decision Comparisons Complete Supplement'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()


def build(root=ROOT):
    root=Path(root);dest=root/'reports/comparators';dest.mkdir(parents=True,exist_ok=True);(dest/'metadata').mkdir(parents=True,exist_ok=True);record=dest/'metadata/appendix_verification.json'
    record.write_text(json.dumps(dict(passed=False,status='building')),encoding='utf-8')
    (dest/(STEM+'.pdf')).unlink(missing_ok=True)
    try:
        inputs={}
        def read(name):
            path=root/name;inputs[name]=sha(path)
            return list(csv.DictReader(path.open(encoding='utf-8',newline='')))
        evaluations=read('results/phase2/comparators/evaluations.csv')
        intervals=read('results/phase2/comparators/intervals.csv')
        for name in ('reporting/overall.py','reporting/edition.py','reports/comparators/comparator_appendix.md','phase2/comparators/protocol.json'):
            inputs[name]=sha(root/name)
        lookup={(r['left_method'],r['rule'],r['cost'],r['fold']):r for r in intervals if r['comparison']=='policy-none'}
        def effect(r):return f"{float(r['loss_delta100']):.2f} [{float(r['ci_low']):.2f}, {float(r['ci_high']):.2f}]"
        items=[dict(kind='title',text=TITLE),dict(kind='paragraph',text='Junjie Yang'),
            dict(kind='paragraph',text='This supplement retains every specified frozen-policy comparison: 720 policy-minus-no-alarm settings and 540 XGBoost-minus-grade settings. Record blocks 1-4 precede supplementary pooled sums. Stage 1 and Stage 2 denote research stages; blocks denote evaluated record groups.'),
            dict(kind='paragraph',text='Negative loss differences favor the left policy. Costs are assumed relative values, not measured money or accident avoidance. All intervals are unadjusted conditional 95% paired moving-block intervals from 2,000 draws of length 32 within blocks; they exclude model fitting, tuning and rule selection. Identical decisions can yield exact zero intervals. No policy is reselected from these tables.'),
            dict(kind='link',text='Integrated manuscript in this edition',target=edition_url('reports/overall/integrated_research_report.pdf')),
            dict(kind='link',text='Comparator protocol and fixed selection rules',target=edition_url('phase2/comparators/protocol.json')),
            dict(kind='link',text='Complete machine-readable evidence and verification',target=edition_url('results/phase2/comparators/README.md'))]
        tables=[]
        for comparison in ('policy-none','ML-grade'):
            items.append(dict(kind='heading',text='Policies minus no alarms' if comparison=='policy-none' else 'Fixed XGBoost minus each grade rule'))
            for cost in (5,10,20,50):
                rows=[]
                if comparison=='policy-none':
                    part=[r for r in evaluations if float(r['cost'])==cost]
                    for r in sorted(part,key=lambda r:(int(r['fold']) if int(r['fold'])>=0 else 4,r['method'],r['rule'])):
                        ci=lookup[(r['method'],r['rule'],r['cost'],r['fold'])]
                        rows.append(['Pooled' if int(r['fold'])<0 else str(int(r['fold'])+1),r['method'],r['rule'],effect(ci)]+[str(int(float(r[k]))) for k in ('tp','fn','fp','alerts')]+[str(int(float(r['excess']))) if r['excess'] else '-'])
                    headers=['Block','Method','Rule','Delta loss / 100 [95% CI]','TP','FN','FP','Alerts','Excess']
                else:
                    part=[r for r in intervals if r['comparison']=='ML-grade' and float(r['cost'])==cost]
                    for r in sorted(part,key=lambda r:(int(r['fold']) if int(r['fold'])>=0 else 4,r['right_method'],r['rule'])):
                        rows.append(['Pooled' if int(r['fold'])<0 else str(int(r['fold'])+1),r['right_method'],r['rule'],effect(r)])
                    headers=['Block','Grade field','Rule','ML minus grade loss / 100 [95% CI]']
                table=dict(kind='table',id=f'appendix_{comparison}_{cost}',caption=f'Table S{len(tables)+1}. '+('Policy minus no alarms' if comparison=='policy-none' else 'XGBoost minus direct grade')+f'; assumed missed-hazard cost r={cost}, false-alert cost 1.',headers=headers,rows=rows)
                tables.append(table);items.append(table)
        write_outputs(items,dest,STEM,TITLE)
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Inches,Pt
        from phase1.generate_report import scrub_report_package
        path=dest/(STEM+'.docx');doc=Document(path)
        for section in doc.sections:
            section.orientation=WD_ORIENT.LANDSCAPE;section.page_width=Inches(11.69);section.page_height=Inches(8.27)
            section.top_margin=section.bottom_margin=Inches(.55);section.left_margin=section.right_margin=Inches(.6)
        for table in doc.tables:
            widths=[.6,1.1,1.5,3.1,.6,.6,.6,.8,.8] if len(table.columns)==9 else [1.,2.,2.,4.7]
            table.autofit=False
            for column,width in zip(table.columns,widths):column.width=Inches(width)
            for row in table.rows:
                for cell,width in zip(row.cells,widths):
                    cell.width=Inches(width)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(2)
        doc.save(path);scrub_report_package(path)
        manifest=dict(edition_tag=EDITION_TAG,generator_sha256=sha(Path(__file__)),inputs=inputs,tables=tables,
            created_at_utc=datetime.now(timezone.utc).isoformat(),scope='Full frozen comparison presentation; no fits, threshold selection or new intervals',
            outputs={name:sha(dest/name) for name in (STEM+'.md',STEM+'.docx')})
        (dest/'metadata/appendix_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
        record.write_text(json.dumps(dict(passed=False,status='awaiting verification')),encoding='utf-8')
        print('Complete decision supplement generated: 8 tables, 1260 settings.')
    except (Exception,KeyboardInterrupt) as error:
        record.write_text(json.dumps(dict(passed=False,status='failed',error_type=type(error).__name__,error=str(error))),encoding='utf-8');raise


if __name__=='__main__':build()
