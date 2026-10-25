"""Build an integrated manuscript from saved, verified two-stage evidence.

Run ``python -m reporting.overall`` in the reporting environment. Model fits,
thresholds, costs and retained experiment outputs are never changed here.
"""
from __future__ import annotations
import csv
import hashlib
import json
import math
from pathlib import Path
import re
from datetime import datetime, timezone

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / 'reports/overall'
STEM = 'integrated_research_report'
TITLE = 'From Forecast Evaluation to Warning Decisions: A Retrospective Study of Seismic Hazard in Underground Mining'


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def match(row, **conditions):
    for key, value in conditions.items():
        actual = row[key]
        if isinstance(value, (float, int)):
            if not actual or float(actual) != value:
                return False
        elif actual != value:
            return False
    return True


def one(rows, **conditions):
    selected = [row for row in rows if match(row, **conditions)]
    if len(selected) != 1:
        raise ValueError(f'Expected one evidence row, found {len(selected)}: {conditions}')
    return selected[0]


def num(value, places=4):
    return f'{float(value):.{places}f}'


def integer(value):
    return str(int(float(value)))


def interval(lo, hi, places=4):
    return f'[{num(lo, places)}, {num(hi, places)}]'


def span(values, places=2):
    values = list(map(float, values))
    lo, hi = min(values), max(values)
    return num(lo, places) if lo == hi else f'{num(lo, places)} to {num(hi, places)}'


class Evidence:
    def __init__(self):
        self.inputs = {}
        self.cache = {}

    def text(self, relative):
        path = ROOT / relative
        self.inputs[relative] = digest(path)
        return path.read_text(encoding='utf-8')

    def csv(self, relative):
        if relative not in self.cache:
            self.cache[relative] = list(csv.DictReader(self.text(relative).splitlines()))
        return self.cache[relative]

    def json(self, relative):
        return json.loads(self.text(relative))

    def image(self, relative):
        self.inputs[relative] = digest(ROOT / relative)
        return relative


def overview_figure(path):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    fig, ax = plt.subplots(figsize=(10, 4.2))
    ax.set(xlim=(0, 10), ylim=(0, 4.2)); ax.axis('off')
    boxes = [(.15, 1.30, 2.0, 1.6, 'Shared records\n\nSource audit\nCommon test rows', '#edf2f6'),
             (3.0, 2.3, 3.25, 1.55, 'Forecast evidence\n\nRanking and probabilities\nFeature use and ablation', '#e5eff9'),
             (3.0, .25, 3.25, 1.55, 'Decision evidence\n\nHistorical A/B/C rules\nCutoff transfer and refitting', '#e7f3ee'),
             (7.1, 1.15, 2.7, 1.9, 'Joint interpretation\n\nLoss and coverage\nDemand and capacity excess', '#f5eee0')]
    for x, y, w, h, label, color in boxes:
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0.03',
                                   facecolor=color, edgecolor='#65717c', linewidth=1))
        ax.text(x+w/2, y+h/2, label, ha='center', va='center', fontsize=13, linespacing=1.5)
    for start, end in [((2.2, 2.5), (2.95, 3.0)), ((2.2, 1.65), (2.95, 1.0)),
                       ((6.3, 3.0), (7.05, 2.6)), ((6.3, 1.0), (7.05, 1.6))]:
        ax.annotate('', xy=end, xytext=start, arrowprops=dict(arrowstyle='->', color='#65717c', lw=1.3))
    fig.tight_layout(pad=.25)
    fig.savefig(path, dpi=320, bbox_inches='tight')
    fig.savefig(Path(path).with_suffix('.svg'), bbox_inches='tight')
    plt.close(fig)


def decision_figure_rows(evaluations, intervals):
    """Select saved frozen-C results; never estimate new intervals or policies."""
    rows=[]
    for fold in (0,1,2,3,-1):
        for budget in (.01,.05,.10,.20):
            r=one(evaluations,method='XGBoost',fold=fold,mechanism='C',cost=10,budget=budget)
            ci=one(intervals,left_method='XGBoost',right_method='none',fold=fold,mechanism='C',cost=10,budget=budget)
            rows.append(dict(model='XGBoost',workflow='fixed',mechanism='C',cost=10,fold=fold,budget=budget,n=integer(r['n']),positives=integer(r['positives']),
                             delta_loss100=ci['loss_delta100'],ci_low=ci['ci_low'],ci_high=ci['ci_high'],
                             alerts=integer(r['alerts']),tp=integer(r['tp']),fn=integer(r['fn']),
                             slots=integer(r['slots']),excess=integer(r['excess'])))
    return rows


def decision_figure(rows, destination):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.ticker import MaxNLocator
    destination=Path(destination)
    with (destination/'decision_tradeoffs.csv').open('w',encoding='utf-8',newline='') as stream:
        writer=csv.DictWriter(stream,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
    fig,axes=plt.subplots(5,2,figsize=(9.2,8.5),gridspec_kw={'width_ratios':[1.25,1]},sharex='col')
    low=min(float(r['ci_low']) for r in rows);high=max(float(r['ci_high']) for r in rows)
    margin=.10*(high-low)
    for index,fold in enumerate((0,1,2,3,-1)):
        part=[r for r in rows if r['fold']==fold]
        ax,counts=axes[index]
        title=('Pooled (supplementary)' if fold<0 else f'Record block {fold+1}')
        ax.set_title(title+f"   n={part[0]['n']}, P={part[0]['positives']}",loc='left',fontsize=12,pad=5)
        for y,r in enumerate(part):
            delta=float(r['delta_loss100']);lo=float(r['ci_low']);hi=float(r['ci_high'])
            ax.plot([lo,hi],[y,y],color='#24668b',lw=1.6)
            ax.plot(delta,y,'o',color='#24668b',ms=4)
            counts.barh(y-.13,int(r['alerts']),height=.24,color='#9dabb6',label='Alerts' if y==0 else None)
            counts.barh(y+.13,int(r['tp']),height=.24,color='#007d70',label='TP' if y==0 else None)
            counts.text(max(int(r['alerts']),int(r['tp']))+.7,y,r['alerts']+' | '+r['tp']+'/'+r['positives'],va='center',fontsize=11,color='#254450')
        ax.axvline(0,color='#59636c',lw=.8,ls='--');ax.set_xlim(low-margin,high+margin)
        ax.set_yticks(range(4),['1%','5%','10%','20%']);ax.set_ylabel('Budget',fontsize=11)
        counts.set_yticks([]);counts.set_xlim(0,max(int(r['alerts']) for r in rows)+18)
        counts.xaxis.set_major_locator(MaxNLocator(integer=True,nbins=5))
        for panel in (ax,counts):
            panel.set_ylim(3.55,-.55)
            panel.spines[['top','right']].set_visible(False)
            panel.tick_params(labelsize=11)
            panel.grid(axis='x',alpha=.13);panel.set_axisbelow(True)
        if fold<0:
            for panel in (ax,counts):panel.set_facecolor('#f0f3f5')
    axes[0,1].set_title('Transferred demand and coverage',fontsize=12,pad=5)
    axes[0,1].legend(frameon=False,loc='upper right',fontsize=11)
    axes[-1,0].set_xlabel('Loss difference vs no alarms / 100 records',fontsize=12)
    axes[-1,1].set_xlabel('Alerts and TP/P (hazardous shifts flagged)',fontsize=12)
    fig.tight_layout(h_pad=1.0,w_pad=2.0)
    fig.savefig(destination/'decision_tradeoffs.png',dpi=320,bbox_inches='tight');plt.close(fig)


def decision_signature(decisions, fold, methods, budgets):
    """Exact row/label/alert tuples, including every comparator at every budget."""
    return tuple((method,tuple((r['row'],r['y'],r[f'C_b{round(100*b):02d}_r10'])
                 for r in sorted(decisions,key=lambda r:int(r['row']))
                 if r['method']==method and (fold<0 or int(r['fold'])==fold)))
                 for method in methods for b in budgets)


def assemble():
    evidence = Evidence()
    s1 = lambda name: evidence.csv('results/phase1/' + name + '.csv')
    s2 = lambda name: evidence.csv('results/phase2/' + name + '.csv')
    source = evidence.json('results/phase1/source_audit.json')
    lock = evidence.json('phase2/locked_plan.json')
    plan = lock['plan']
    # Record source-certification scope without rewriting the source records.
    certificates = {}
    for relative in ['results/phase2/verification.json', 'results/phase2/supplement_verification.json',
                     'results/phase2/comparators/verification.json']:
        record = json.loads((ROOT / relative).read_text(encoding='utf-8'))
        if not record.get('passed'):
            raise RuntimeError('Passed source verification required: ' + relative)
        if relative=='results/phase2/verification.json' and not record.get('replay'):
            raise RuntimeError('A retained full Stage 2 replay is required before authoring')
        certificates[relative] = dict(sha256=digest(ROOT / relative), scope=record.get('scope'),
                                      replay=record.get('replay'), checks=record.get('checks'))
    old1 = evidence.text('reports/phase1/technical_note.md')
    old2 = evidence.text('reports/phase2/phase2_threshold_transfer_report.md')
    evidence.text('CONTRIBUTIONS.md')
    ref1 = {int(n): text for n, text in re.findall(r'^\[(\d+)\] (.+)$', old1, re.M)}
    ref2 = {int(n): text for n, text in re.findall(r'^\[(\d+)\] (.+)$', old2, re.M)}
    references = {f's1_{n}': text for n, text in ref1.items()}
    references.update(kunsch=ref2[4], warning=ref2[12])
    references.update(
        naeini='Pakdaman Naeini M, Cooper G, Hauskrecht M. Obtaining Well Calibrated Probabilities Using Bayesian Binning. Proceedings of the AAAI Conference on Artificial Intelligence, 2015, 29(1). DOI: 10.1609/aaai.v29i1.9602.',
        guo='Guo C, Pleiss G, Sun Y, Weinberger KQ. On Calibration of Modern Neural Networks. Proceedings of ICML, PMLR 70, 2017:1321-1330. https://proceedings.mlr.press/v70/guo17a.html',
        brier='Brier GW. Verification of Forecasts Expressed in Terms of Probability. Monthly Weather Review, 1950, 78(1):1-3. DOI: 10.1175/1520-0493(1950)078<0001:VOFEIT>2.0.CO;2.')
    used = []

    def cite(*keys):
        for key in keys:
            if key not in used:
                used.append(key)
        return '[' + ','.join(str(used.index(key)+1) for key in keys) + ']'

    models = ('LR', 'CART', 'XGBoost')
    ablation = s1('feature_ablation'); ab_ci = s1('feature_ablation_intervals')
    full = {m: one(ablation, model=m, variant='full') for m in models}
    n, positives = int(full['LR']['n']), int(full['LR']['positives'])
    prevalence = positives/n
    shared = s1('shared_test_comparison'); protocol_ci = s1('paired_uncertainty')
    xrandom = s1('xgboost_shared_by_seed'); xci = s1('xgboost_protocol_interval')[0]
    gaps = s1('gap_attenuation_all_models')
    if {r['model'] for r in gaps}!=set(models):
        raise ValueError('Run phase1.gap_summary before reporting: all three gap rows required')
    calibration = s1('calibration')
    raw = next(r for r in calibration if r['model']=='LR' and r['variant'].startswith('as'))
    corrected = next(r for r in calibration if r['model']=='LR' and r['variant'].startswith('after'))
    probability = s1('probability_reference_comparison')
    historical = {m: one(probability, model=m, validation='time', fold=-1) for m in models}
    removed = one(ablation, model='XGBoost', variant='without_seismic_activity')
    removed_ci = one(ab_ci, model='XGBoost', variant='without_seismic_activity')
    importance = sorted([r for r in s1('shap_importance') if match(r, validation='time', fold=-1)],
                        key=lambda r: -float(r['mean_abs']))
    folds = evidence.json('results/phase2/fold_manifest.json')
    folds = sorted([m for m in folds if m['model']=='XGBoost' and m['scheme']=='time'], key=lambda m:m['fold'])
    phase = s1('feature_ablation_by_fold')
    evaluations = s2('evaluations'); pooled = s2('pooled_evaluations')
    main = [r for r in pooled if match(r, model='XGBoost', workflow='fixed', cost=10)]
    selected_c = sorted([r for r in main if r['mechanism']=='C'], key=lambda r:float(r['budget']))
    audits = s2('threshold_explanations'); all_audits = s2('audits'); transfers = s2('refit_transfer')
    comparisons = s2('comparisons'); capacity = s2('capacity_tradeoffs')
    boundary = s2('frozen_policy_cost_boundaries'); scenarios = s2('prevalence_decision_value')
    supplement = s2('comparators/evaluations'); supplement_ci = s2('comparators/intervals')
    supplement_decisions = s2('comparators/decisions')
    grade_metadata = evidence.json('results/phase2/comparators/grade_metadata.json')
    evidence.json('phase2/comparators/protocol.json')
    xv = [r for r in evaluations if match(r, model='XGBoost', scheme='time', workflow='fixed')]
    va = [r for r in xv if match(r, mechanism='A', cost=10)]
    vc = [r for r in xv if r['mechanism']=='C']
    silent = [r for r in audits if match(r, model='XGBoost', scheme='time', reference_alerts=0)]
    causes = {label:sum(r['no_alarm_reason']==label for r in silent)
              for label in ('cost_selects_no_alarm','capacity_changes_to_no_alarm')}
    over_a, over_c = sum(float(r['excess'])>0 for r in va), sum(float(r['excess'])>0 for r in vc)
    reductions = ', '.join(f'{float(r["gap_reduction_fraction"])*100:.1f}% for {r["model"]}' for r in gaps)
    a20=one(main,mechanism='A',budget=.2);c20=one(main,mechanism='C',budget=.2)
    grade_b={field:sorted([r for r in supplement_decisions if r['method']==field],key=lambda r:int(r['row']))
             for field in ('seismic','seismoacoustic','ghazard')}
    for field in ('seismoacoustic','ghazard'):
        if [(r['row'],r['B_r10']) for r in grade_b[field]]!=[(r['row'],r['B_r10']) for r in grade_b['seismic']]:
            raise ValueError('The stated common grade-B decision no longer holds')
    if not all((r['B_r10']=='True')==(int(r['fold'])==0) for r in grade_b['seismic']):
        raise ValueError('The stated phase pattern for grade B no longer holds')
    gb=one(supplement,method='seismic',fold=-1,mechanism='B',cost=10)
    xb=one(supplement,method='XGBoost',fold=-1,mechanism='B',cost=10)
    fewer_fp=integer(float(gb['fp'])-float(xb['fp']));fewer_tp=integer(float(gb['tp'])-float(xb['tp']))
    information=evidence.csv('results/phase2/information/evaluations.csv')
    information_ci=evidence.csv('results/phase2/information/intervals.csv')
    evidence.csv('results/phase2/information/rows.csv')
    evidence.json('phase2/information/protocol.json')
    evidence.json('results/phase2/information/run_manifest.json')
    info_check=evidence.json('results/phase2/information/verification.json')
    if not info_check.get('passed') or not info_check.get('refit'):
        raise RuntimeError('Current matched-information calculations and refit replay required')
    evidence.text('phase2/information/verify.py')
    evidence.text('phase2/information/provenance.py')
    evidence.inputs['results/phase2/information/historical_context.zip']=digest(ROOT/'results/phase2/information/historical_context.zip')
    certificates['results/phase2/information/verification.json']=dict(
        sha256=digest(ROOT/'results/phase2/information/verification.json'),scope=info_check['scope'],refit=True,checks=info_check['checks'])
    information20=one(information,fold=-1,budget=.2)
    information20_ci=one(information_ci,fold=-1,budget=.2,block_length=32)
    from reporting.visuals import supporting_figures, monitoring_figure
    supporting_figures(s1('calibration_bins'),s1('shap_importance'),DEST)
    monitoring_figure(information,information_ci,DEST)
    evidence.text('reporting/visuals.py');evidence.text('reporting/word_math.py')

    def effect(row):
        return num(row['loss_delta100'],2)+' '+interval(row['ci_low'],row['ci_high'],2)

    items, tables = [], []

    def para(text): items.append(dict(kind='paragraph', text=text))
    def equation(key): items.append(dict(kind='equation',id=key))
    def heading(text, page=False):
        if page and text=='References':
            items.append(dict(kind='break'))
        items.append(dict(kind='heading', text=text))
    def table(key, caption, headers, rows, part='', continuation=False):
        number=sum(not t.get('continuation',False) for t in tables)+(not continuation)
        title = f'Table {number}{part}. {caption}'
        value = dict(id=key, caption=title, continuation=continuation, headers=headers, rows=[[str(v) for v in row] for row in rows])
        tables.append(value); items.append(dict(kind='table', **value))
    def figure(relative, caption, width=6.4):
        evidence.image(relative)
        items.append(dict(kind='figure', path=relative, caption=caption, width=width))

    items.append(dict(kind='title', text=TITLE))
    para('Junjie Yang'); para('Mining Engineering, Fuzhou University')
    heading('Abstract')
    para(f'Seismic warning rules can change their loss, coverage and workload on later records. We examine this transfer using complementary forecast and decision analyses on {n:,} common UCI Seismic Bumps test records with {positives} hazardous-shift labels. Matching test records substantially narrows random-versus-record-order average-precision differences, showing cohort sensitivity without isolating leakage or drift. We then select budget-only, cost-only and capacity-constrained cost rules on historical reference data and evaluate their frozen transfer. '
         f'For fixed XGBoost models at an assumed missed-hazard cost ten times the false-alert cost, capacity-constrained cost rules show pooled loss differences of {span(float(r["loss_delta_no_alarm100"]) for r in selected_c)} units per 100 records against no alarms; intervals for the nonzero differences include zero. At a 20% historical budget, these rules produce {integer(c20["alerts"])} alerts and {integer(c20["tp"])} detections, compared with {integer(a20["alerts"])} alerts and {integer(a20["tp"])} detections for budget-only rules, among {positives} hazardous shifts. Historical capacity feasibility does not ensure later feasibility, and refitting can change decisions at an unchanged cutoff. '
         'A separately specified exploratory comparison at matched grade allocations and test-batch slots does not establish a consistent detection increment from adding monitoring variables beyond grades and shift type. The findings support joint evaluation of loss, coverage and workload on overlapping records; recorded order remains a temporal proxy.')
    para('Keywords: seismic hazard; forecast evaluation; average precision; warning decisions; threshold transfer; inspection capacity')
    heading('1 Introduction and research objective')
    para('Mine-monitoring scores become engineering decisions when a rule selects subsequent shifts for inspection. Historical loss minimization or capacity feasibility does not determine what the same rule will demand or detect on later records. This study asks: how do historically selected seismic warning rules perform on subsequent records in loss, hazardous-shift coverage and alarm workload, and what forecast-evaluation evidence supports interpreting those outcomes?')
    para('Two questions organize the analysis. First, how does evaluation on common test records change the apparent forecast-protocol gap? Second, how do loss, coverage and workload change when historical warning rules transfer under assumed costs and capacity? Calibration, TreeSHAP and refitted ablations provide supporting evidence about probabilities and model use of information. The earlier within-grade ranking pilot remains separate. A subsequent matched-information analysis supports the second question by fixing test-batch slots and grade allocations; its core result is included here, with full evidence in the exploratory supplement. The two stages retain their own fitting and threshold protocols; one is not a prospective validation of the other.')

    heading('1.1 Relation to mining and decision research', page=True)
    para(f'The UCI task concerns a high-energy seismic bump in the next shift, rather than a confirmed rockburst accident {cite("s1_1")}. Earlier rule-induction work used working-face-specific monitoring and examined rule transfer {cite("s1_9")}. '
         f'Janusz and colleagues compared expert and sensor information and studied cold start with identified mining sites and structured histories {cite("s1_12")}. Those cohorts and forecast targets do not supply a directly comparable performance leaderboard for the present shift-summary mirror. '
         f'The microseismic monitoring review distinguishes waveform detection and location from forecasting {cite("s1_13")}. This study evaluates supplied shift summaries, not event detection from waveforms.')
    para(f'Threshold-choice theory relates scores, operating conditions and expected classification loss {cite("s1_14")}. Risk-informed warning identification connects warnings to assessment and response, but its assessment-based Seismic Bumps comparison differs from this study\'s subsequent-shift event label {cite("warning")}. '
         f'Precision-recall interpretation also depends on the evaluation prevalence {cite("s1_15")}. These sources motivate explicit comparators; they do not establish that a retrospectively selected rule will work at a new mine or future operating condition.')
    para(f'The study contributes three connected elements. First, common-cohort forecast assessment quantifies how the apparent AP protocol gap changes with the evaluated records: the gaps shrink by {reductions}, without identifying a causal share of leakage or drift. Second, the decision analysis jointly measures transferred loss, detections, workload and capacity excess, explains historical no-alarm choices and compares fixed with refitted models. A matched-allocation supplement then asks whether monitoring information adds coverage beyond grades and shift type. Third, the implementation connects source row mappings, saved predictions, historical threshold candidates and row-level decisions to verification through separate calculation paths. This is an application and empirical evaluation of established learning, explanation, threshold-choice and resampling methods.')
    figure('reports/overall/evidence_overview.png', 'Figure 1. Conceptual organization of the study. Shared records support complementary forecast and decision analyses. The branches do not imply that corrected Stage 1 probabilities are inputs to Stage 2 rules.')

    heading('2 Data, target and common evaluation records', page=True)
    para(f'The official source contains {source["original_rows"]:,} records and {source["original_positives"]} positive shift labels. An ARFF-to-mirror audit reproduces first-occurrence deduplication without reordering: {source["removed_rows"]} exact duplicate occurrences, all negative, are removed to yield {source["mirror_rows"]:,} mirror records {cite("s1_1","s1_2")}. '
         'This verifies the mirror transformation; it does not establish that identical monitoring records could not have represented distinct real shifts. A positive label denotes a bump above 10,000 J in the subsequent eight-hour shift. It is not a catalogue of independent accidents or a measured intervention benefit.')
    para('The retained 17-column design contains three hazard assessments (seismic, seismoacoustic, ghazard), four geophone variables, nine seismic-activity variables and shift type. The geophone deviations already use historical summaries supplied with the data. The total count nbumps is omitted from this design; constant energy-band columns remain. Neither timestamps nor working-face identifiers are available to validate cross-row continuity. Recorded order is therefore the stated temporal proxy, not a verified sequence of consecutive shifts from one working face.')
    rows=[]
    for m in folds:
        f=m['fold']; row=one(phase, model='XGBoost', variant='full', fold=f)
        ref=one(audits, model='XGBoost', scheme='time', fold=f, rule='C_b01_r05')
        rows.append([f+1, row['n'], row['positives'], f'{100*float(row["prevalence"]):.2f}%',
                     ref['reference_n'], ref['reference_positives']])
    table('cohorts', 'Shared test census and Stage 2 reference areas. Phases are record blocks, not identified mining sites.',
          ['Phase','Test n','Test +','Test prevalence','Stage 2 ref n','Stage 2 ref +'], rows)
    para(f'The common test cohort has {positives}/{n} positives ({100*prevalence:.2f}%), distinct from the full-mirror prevalence of {100*source["original_positives"]/source["mirror_rows"]:.2f}%. '
         'The 774-record holdout is wholly within the common test rows and uses a different training cutoff. It was already inspected and is supplementary, overlapping evidence; it is neither an additional independent cohort nor an untouched confirmation set. Sample counts are not added across stages or protocols.')

    heading('3 Two complementary evaluation protocols', page=True)
    para(f'LR uses balanced class weights and an L2 penalty of 1. Bagged CART uses 60 unweighted bootstrap trees, depth 6 and minimum leaf size 20. XGBoost uses binary logistic loss, histogram trees, learning rate 0.05, minimum child weight 10 and L2 penalty 5 {cite("s1_6")}. '
         'The grid is depths 2/3 crossed with 80/160 rounds, with seed 7, two threads and full row/column sampling. Tuning occurs within the supplied training area. Models have deliberately asymmetric weighting and tuning; these are comparisons of retained workflows, not equal-budget algorithm searches.')
    table('protocols', 'Workflow definitions. Numerical cutoffs and alert counts are comparable only under matched workflows.',
          ['Item','Forecast study (Stage 1)','Decision study (Stage 2)'],[
              ['Main comparison','Random versus expanding record order on common test rows','Historical A/B/C rules transferred to the same later blocks'],
              ['Model fitting','Outer models refit on supplied training rows','Primary fixed model fits first 80% of outer history'],
              ['Policy reference','Last 20% of history; separate earlier reference model','Last 20% of history; scores from primary fixed model'],
              ['XGBoost choice','Separate training-local searches for outer/reference models','Search inside fitting area; same chosen settings for refitted control'],
              ['Budget cutoff','Exclude boundary score and all ties','Canonical whole-score groups; alerts score >= cutoff'],
              ['Primary transfer','Reference cutoff applied after outer refitting','Freeze model and cutoff; refit on all history as a control'],
              ['Probability scale','LR prior correction for probability evaluation; raw scores for alerts','Raw scores for fixed and refitted alert rules']])
    para(f'Stage 1 uses ten stratified random folds and five seeds; random predictions are restricted to the same {n:,} rows as record-order evaluation. Random AP is the mean of seed-specific AP values. Expanding-window fits train on earlier record blocks. '
         f'Structured validation makes the tested population and information boundary explicit {cite("s1_3","s1_4")}. Matching test rows does not equalize training sizes or remove all protocol differences.')
    para('In Stage 2, the first 80% of each historical prefix is the fitting area and the remainder is the policy reference. XGBoost chooses parameters on a further 80/20 split inside the fitting area using AP, with negative Brier score if the inner reference has no positives. A separate refitted control adds the policy-reference rows to fitting, retains the selected parameters and applies the identical numerical threshold. Test labels enter evaluation, not model, parameter or cutoff selection. The Stage 2 plan was locked before its fits, after the dataset had already been examined in Stage 1; the complete study remains exploratory and retrospective.')

    heading('3.1 Probability, explanation and decision definitions', page=True)
    para(f'AP is the sum of precision values weighted by recall increments, with tied scores evaluated as a group: AP = sum_k P_k (R_k - R_(k-1)). The horizontal PR reference is the evaluation cohort prevalence pi {cite("s1_15")}. AP-minus-pi is reported as a descriptive difference, not a permutation test or an interval estimate. '
         'It neither removes all differences between cohorts nor proves decision value. LR prior correction uses only each outer training fold\'s unweighted prevalence; CART and XGBoost use their fitted outputs.')
    para(f'Expected calibration error is the bin-count-weighted absolute difference between mean event probability and observed hazardous fraction {cite("naeini","guo")}. The present binary event-probability definition uses up to ten quantile bins, collapses duplicate edges and keeps tied probabilities together. Guo and colleagues instead summarize predicted-class confidence and accuracy using equally spaced bins. Binning is an assessment step, not a calibration model fitted on test labels. Brier score (BS) is the binary event-probability mean squared error {cite("brier")}; the named historical training-prior or retrospective test-prevalence predictor supplies its reference; skill is undefined when the reference error is zero.')
    equation('brier_skill')
    para(f'Prior-only distribution change is a specific modeling assumption {cite("s1_10")}. A probability correction is monotone within a fold; differing fold-specific corrections should not be assumed to preserve pooled rankings. '
         f'Native exact TreeSHAP decomposes each XGBoost raw log-odds prediction into feature contributions and a bias term, using training leaf covers as its tree-path-dependent reference {cite("s1_7","s1_8")}. Refitted group ablations ask a different question: how predictions change after a model is trained without the group. They are not causal feature effects.')
    para(f'Cost-sensitive decisions use the hypothetical loss L = r FN + FP, with false-alert cost 1 {cite("s1_11")}. Historical budgets are 1%, 5%, 10% and 20%; r takes 5, 10, 20 and 50. '
         'A maximizes whole-group alerts within floor(B Nref). B minimizes reference loss. C minimizes it among capacity-feasible groups. Ties in B/C loss select fewer alerts. Equal scores move together; nonempty partial sets use the lowest included score. No alarms use +infinity and all alarms use -infinity. These conventions can underfill capacity and do not impose later capacity.')
    para('An alert represents one candidate shift for inspection. TP is a hazardous shift flagged, FP a non-hazardous shift flagged and FN a hazardous shift missed. In the following equations, r is the assumed missed-hazard cost relative to false-alert cost 1, N is the evaluated record count, B is the historical capacity fraction and Ntest is the current test-block size. Loss and its difference from no alarms are:')
    equation('loss100'); equation('delta_none100')
    para('Negative loss differences mean lower hypothetical loss. Capacity excess is calculated separately within each test block:')
    equation('excess')
    para('All alerts remain in evaluation, including excess demand. No dispatch, truncation, queueing, separate true-positive inspection cost or accident-severity cost is modeled. Pooled excess sums phase excesses, without cancellation by spare slots elsewhere.')
    para(f'Paired moving-block intervals use 2,000 replicates and block length 32, separately within test phases, without wrapping across endpoints or phase boundaries {cite("s1_5","kunsch")}. '
         'Both compared methods use the same sampled rows; pooled sampling preserves phase sizes. Intervals condition on saved predictions, fitted rules and the current cohort, and exclude full fitting/tuning/selection uncertainty. The original contrasts cover C-A, C-B, secondary B-A and refitted-fixed. The separately specified supplement adds policy-minus-no-alarm and ML-minus-direct-grade contrasts; its rules, seeds and results are recorded separately. Intervals are unadjusted exploratory comparisons.')
    para(f'The direct-rule supplement thresholds seismic, seismoacoustic and ghazard separately on the original historical reference rows {cite("s1_1")}. UCI explicitly defines seismic severity from a to d; all three ARFF domains use these codes. The common a<b<c<d ordering is the study\'s ordinal interpretation for the two other assessments, not a verified deployed policy. '
         'The mirror contains seismic codes '+', '.join(grade_metadata['observed_counts']['seismic'])+
         '; the two other fields contain a, b and c, with no observed d. Each field uses the same A/B/C objectives, whole-grade ties and canonical thresholds as ML. Duplicate reference alert sets are collapsed: none uses +infinity, all uses -infinity, and partial sets use their minimum included observed grade. No best grade field is selected from test results. Saved ML scores, cutoffs and alerts are reused exactly; no model is retrained.')

    heading('4 Forecast evidence on comparable test records', page=True)
    rows=[]
    for model in models:
        if model=='XGBoost':
            random_ap=sum(float(r['random_ap']) for r in xrandom)/len(xrandom); ci=xci
        else:
            record=next(r for r in shared if ('logistic' in r['model'])==(model=='LR'))
            random_ap=float(record['random_pr_auc']); ci=one(protocol_ci, model=model, block_length=32)
        ap=float(full[model]['pr_auc'])
        rows.append([model,num(random_ap),num(ap),num(ap-prevalence),num(ci['delta']),interval(ci['ci_low'],ci['ci_high'])])
    table('ranking', f'AP on {n:,} common test records; prevalence reference pi = {num(prevalence)}. AP-minus-pi concerns the record-order column.',
          ['Model','Random AP','Record-order AP','AP minus pi','Random minus order','95% paired interval'], rows, part='a')
    gaprows=[[r['model'],num(r['original_gap'],5),num(r['same_test_gap'],5),
              f"{100*(1-float(r['same_test_gap'])/float(r['original_gap'])):.1f}%"] for r in gaps]
    table('gap_attenuation','Attenuation of the AP protocol gap after matching test records.',
          ['Model','Different-cohort gap','Common-cohort gap','Attenuation'],gaprows,part='b',continuation=True)
    para(f'The different-cohort gap compares random validation on {source["mirror_rows"]:,} records with record-order validation on {n:,}; the common-cohort gap evaluates both on the same {n:,} records. Attenuation = 1 - common-cohort gap / different-cohort gap, calculated before rounding. The record-order predictions and fitted random models are unchanged. Attenuation is descriptive cohort sensitivity, not a causal decomposition or a confidence interval.')
    para('The interval in Table 3a belongs to the protocol difference, not to AP-minus-pi. All common-cohort point AP values exceed the prevalence reference; that descriptive observation alone is not a significance test against a random-ranking null. Finite-sample random-order AP and the horizontal PR reference are distinct concepts. No global label permutation is introduced.')
    para(f'Restricting the comparison to common rows reduces the original different-cohort gap by {reductions}. The comparison changes which records contribute to the random score, not its fitted models. Cohort matching establishes evaluation-population sensitivity, while residual differences combine training history, sample size and tuning. Neither attenuation nor the residual is a causal estimate of leakage or drift. Table 3b makes the original and matched gaps directly comparable; the Stage 1 technical report retains the full protocol detail.')
    phase_sentences=[]
    for m in folds:
        r=one(phase, model='XGBoost', variant='full', fold=m['fold'])
        phase_sentences.append(f'phase {m["fold"]+1}: AP {num(r["pr_auc"])} at prevalence {num(r["prevalence"])}')
    para('Stage-specific XGBoost results further separate ranking from cohort composition: ' + '; '.join(phase_sentences) + '. These are Stage 1 outer-model predictions, not Stage 2 fixed-model scores. The supplementary holdout overlaps these records and is not used to choose a model, subset or budget.')

    heading('4.1 Probability improvement depends on the reference', page=True)
    para(f'Training-prior correction reduces LR ECE from {num(raw["ece"])} to {num(corrected["ece"])} and Brier score from {num(raw["brier"])} to {num(corrected["brier"])}. '
         'This compares the same model outputs before and after the probability-scale correction. Brier skill asks whether the corrected forecasts improve on a separately defined reference, so the sign of skill can differ across references without contradicting improved ECE.')
    table('probability', 'Common-cohort probability references. LR is prior-corrected; CART and XGBoost retain raw probabilities.',
          ['Model','Model Brier','Historical Brier','Historical skill','Retrospective Brier','Retrospective skill'],
          [[m,num(historical[m]['model_brier']),num(historical[m]['historical_brier']),num(historical[m]['historical_skill']),
            num(historical[m]['oracle_brier']),num(historical[m]['oracle_skill'])] for m in models])
    figure('reports/overall/calibration_detail.png', 'Figure 2. LR calibration before and after training-prior correction on the common test cohort. Quantile bins retain ties; point size reflects bin count. The fitted panel spans 0-100%; the corrected panel zooms to 0-30% on both axes. Bin counts and n are displayed.', width=6.2)
    para('The historical constant is deployable within each recorded protocol because it uses training labels only. The retrospective constant uses the evaluated prevalence for assessment and does not generate predictions. Stage 2 does not apply cost cutoffs to corrected LR probabilities: its A/B/C rules operate on raw reference scores. A calibration gain therefore does not constitute an observed reduction in Stage 2 warning loss.')

    heading('4.2 Feature use and incremental predictive evidence', page=True)
    rows=[]
    for model in models:
        rating=one(ablation, model=model, variant='ratings_only'); ci=one(ab_ci, model=model, variant='ratings_only')
        rows.append([model,num(full[model]['pr_auc']),num(rating['pr_auc']),
                     num(float(full[model]['pr_auc'])-float(rating['pr_auc'])),interval(-float(ci['ci_high']),-float(ci['ci_low']))])
    table('ablation', 'Full design versus a refitted learner using all three hazard-rating fields.',
          ['Model','Full AP','Ratings-only AP','Full minus ratings','95% paired interval'], rows)
    para('Ratings-only denotes a separately fitted LR, CART or XGBoost using seismic, seismoacoustic and ghazard together. It is not a direct grade-to-alarm rule or the mine\'s deployed practice. The paired gain intervals include zero under the retained configurations. XGBoost ablations reuse full-design training-selected parameters, so the contrasts are conditional design comparisons rather than equally retuned optima for every subset.')
    figure('reports/overall/shap_readable.png', 'Figure 3. Full-feature XGBoost attribution and phase-specific shares. Contributions describe the fitted models on raw log-odds; the eight displayed features are selected by global mean absolute contribution. Each heatmap cell is its share of the sum over all 17 features in the block.')
    para(f'The leading contributions are {", ".join(r["feature"] for r in importance[:3])}. Removing the seismic-activity group and refitting yields AP {num(removed["pr_auc"])} versus {num(full["XGBoost"]["pr_auc"])} for the full design; '
         f'the paired difference is {num(removed_ci["delta"])} with interval {interval(removed_ci["ci_low"], removed_ci["ci_high"])}. The interval includes zero. Large attribution and uncertain refitted benefit can coexist because they concern different models and questions. Correlated inputs, interactions and sampling variability are possible explanations, not identified causes. No subset is selected from these test results.')

    heading('5 Warning loss, detections and inspection demand', page=True)
    mainrows=[]
    ordered=[one(main, mechanism='none')]
    ordered += sorted([r for r in main if r['mechanism']=='A'], key=lambda r:float(r['budget']))
    ordered += [one(main, mechanism='B')]+selected_c
    for r in ordered:
        mainrows.append([r['mechanism'], '-' if not r['budget'] else f'{100*float(r["budget"]):.0f}%',
                         num(r['loss100'],2),num(r['loss_delta_no_alarm100'],2),integer(r['tp']),integer(r['fn']),
                         integer(r['fp']),integer(r['alerts']),'-' if not r['excess'] else integer(r['excess'])])
    table('decision', f'Stage 2 fixed XGBoost, r=10, pooled over {n:,} records. Negative loss differences favor the policy over no alarms.',
          ['Rule','Budget','L100','Delta vs none','TP','FN','FP','Alerts','Excess'], mainrows)
    para('The point loss of C equals the no-alarm comparator at 1% and is higher at the other three budgets. The new frozen-policy intervals below assess this particular contrast. Lower C loss than A is compatible with a higher loss than no alarms: reducing false alerts can lower assumed cost while also reducing detections.')
    descriptions=[]
    for f in range(4):
        part=[r for r in xv if match(r, fold=f, mechanism='C', cost=10)]
        descriptions.append(f'phase {f+1}: {span(r["loss_delta_no_alarm100"] for r in part)}')
    para('At r=10 the phase-specific C-minus-no-alarm loss ranges are ' + '; '.join(descriptions) + ' units per 100 records. These ranges retain all four budgets; pooled findings are not attributed uniformly to every phase. Historical budgets constrain reference selection, so methods with the same budget can generate different later workloads.')
    pair=one(capacity, model='XGBoost', scheme='time', cohort='pooled', workflow='fixed', cost=10, budget=.2)
    para(f'For the 20% budget at r=10, pooled C-A changes detections from {integer(pair["A_tp"])} to {integer(pair["C_tp"])} and reduces excess demand by {integer(-float(pair["excess_delta"]))} alerts. '
         f'The loss difference is {num(pair["loss100_delta"],2)} per 100 records. This is a measured workload-coverage tradeoff under the assumed loss, not a demonstrated safety improvement. Actual inspection execution and accident avoidance are not observed.')

    for f in (0,1):
        ci=one(comparisons,model='XGBoost',comparison='C-A',workflow='fixed',budget=.2,cost=10,fold=f)
        para(f'At r=10 and budget 20%, phase {f+1} C-A is {num(ci["loss_delta100"],2)} per 100 records, with paired 95% interval {interval(ci["ci_low"],ci["ci_high"],2)}. '
             + ('The interval contains zero.' if float(ci['ci_low'])<=0<=float(ci['ci_high']) else 'The interval is entirely below zero.'))

    heading('5.1 Frozen policies compared with no alarms')
    decision_figure(decision_figure_rows(supplement,supplement_ci),DEST)
    evidence.text('reports/overall/decision_tradeoffs.csv')
    figure('reports/overall/decision_tradeoffs.png', 'Figure 4. Fixed-XGBoost C at r=10 across four historical budgets. Left: loss versus no alarms per 100 records (negative favors C), with conditional 95% block intervals. Right: alerts (gray) and hazardous shifts flagged, TP (green); P denotes all hazardous shifts. Shaded pooling is supplementary. Zero intervals reflect identical decisions. Intervals exclude fitting and selection uncertainty.',width=6.4)
    rows=[]
    for f in (0,1,2,3,-1):
        groups={}
        for b in plan['budgets']:
            groups.setdefault(decision_signature(supplement_decisions,f,('XGBoost',),(b,)),[]).append(b)
        for budgets in groups.values():
            b=budgets[0]
            r=one(supplement,method='XGBoost',fold=f,mechanism='C',cost=10,budget=b)
            ci=one(supplement_ci,left_method='XGBoost',right_method='none',fold=f,mechanism='C',cost=10,budget=b)
            rows.append(['Pooled' if f<0 else str(f+1),', '.join(f'{100*v:.0f}%' for v in budgets),effect(ci),integer(r['tp']),integer(r['fn']),integer(r['alerts'])])
    table('no_alarm_intervals','Frozen XGBoost C minus no alarms at r=10, per 100 records. Brackets are paired 95% intervals; negative favors C.',
          ['Phase','Budget','Delta [95% interval]','TP','FN','Alerts'],rows)
    para('Comma-separated budgets share exactly the same aligned row-level decisions within that phase, not merely rounded loss or interval endpoints. The complete appendix retains each setting separately. Phase 1 C at 5% and 10% gives positive loss intervals and no detections; the nonzero pooled intervals cross zero. Phase and pooled evidence therefore support different loss-direction assessments. Where C produces no alerts, every row-level difference is zero and the interval is [0,0] by construction. This establishes identical observed decisions, not reliable transfer. No alarms detect no hazardous shifts and remain a loss comparator rather than an operating recommendation.')

    heading('5.2 Direct hazard-grade comparators')
    rows=[]
    for method in ('XGBoost','seismic','seismoacoustic','ghazard'):
        groups={}
        for b in plan['budgets']:
            groups.setdefault(decision_signature(supplement_decisions,-1,(method,),(b,)),[]).append(b)
        for budgets in groups.values():
            b=budgets[0]
            r=one(supplement,method=method,fold=-1,mechanism='C',cost=10,budget=b)
            ci=one(supplement_ci,left_method=method,right_method='none',fold=-1,mechanism='C',cost=10,budget=b)
            excess=[integer(one(supplement,method=method,fold=-1,mechanism='C',cost=10,budget=v)['excess']) for v in budgets]
            rows.append([', '.join(f'{100*v:.0f}%' for v in budgets),method,effect(ci),integer(r['tp']),integer(r['fn']),integer(r['fp']),integer(r['alerts']),', '.join(excess)])
    table('grade_decisions','C rules at r=10 on the same pooled records. Reference selection is separate for each method; future workloads need not match.',
          ['Budget','Method','Delta vs none [95%]','TP','FN','FP','Alerts','Excess'],rows)
    para('Budgets are grouped only when their full transferred decision vectors match. Comma-separated excess counts follow the listed budget order: identical alarms can exceed different numbers of capacity slots. The appendix retains the ungrouped phase and pooled settings.')
    grade=one(supplement,method='ghazard',fold=-1,mechanism='C',cost=10,budget=.2)
    grade_ci=one(supplement_ci,left_method='ghazard',right_method='none',fold=-1,mechanism='C',cost=10,budget=.2)
    para(f'At r=10, seismic and seismoacoustic C select no alarms across the displayed budgets. The ghazard rules yield {integer(grade["alerts"])} alerts, {integer(grade["tp"])} detections and {integer(grade["fp"])} false alerts, with loss difference {effect(grade_ci)} versus no alarms. '
         'Their repeated pooled counts reflect the same transferred decisions, not independent replications. The negative point difference does not establish loss improvement because its interval crosses zero. Coarse reference grade groups can underfill budgets; later frequencies can also exceed them.')
    rows=[]
    for f in (0,1,2,3,-1):
        groups={}
        for b in plan['budgets']:
            groups.setdefault(decision_signature(supplement_decisions,f,('XGBoost','seismic','seismoacoustic','ghazard'),(b,)),[]).append(b)
        for budgets in groups.values():
            b=budgets[0]
            cells=['Pooled' if f<0 else str(f+1),', '.join(f'{100*v:.0f}%' for v in budgets)]
            for field in ('seismic','seismoacoustic','ghazard'):
                ci=one(supplement_ci,left_method='XGBoost',right_method=field,fold=f,mechanism='C',cost=10,budget=b)
                cells.append(effect(ci))
            rows.append(cells)
    table('grade_intervals','Frozen XGBoost minus each historically selected grade C rule at r=10. Brackets are paired 95% intervals; negative favors ML.',
          ['Phase','Budget','ML minus seismic','ML minus seismoacoustic','ML minus ghazard'],rows)
    para('All nondegenerate pooled C-minus-grade contrasts at the displayed cost include zero; identical decisions give exact zero contrasts. These C comparisons do not establish additional ML loss value over the direct grade rules. They also do not prove equivalence: the intervals retain both loss improvement and loss increase. The appendix retains A/B/C for every planned cost and budget, including per-phase and pooled detections, workload, excess and paired contrasts. The direct rules differ from the jointly fitted ratings-only models in Table 5.')
    b_grade=[one(supplement_ci,left_method='XGBoost',right_method=field,fold=-1,mechanism='B',cost=10) for field in ('seismic','seismoacoustic','ghazard')]
    b_none=one(supplement_ci,left_method='XGBoost',right_method='none',fold=-1,mechanism='B',cost=10)
    b_values=[effect(r) for r in b_grade]
    b_text=b_values[0]+' for each grade field' if len(set(b_values))==1 else '; '.join(field+': '+value for field,value in zip(('seismic','seismoacoustic','ghazard'),b_values))
    para(f'The selection mechanism changes the comparison. For cost-only B at r=10, pooled ML-minus-grade loss is {b_text}, with all three intervals below zero. In contrast, B-minus-no-alarm is {effect(b_none)}. '
         f'The three grade B rules make exactly the same row-level decisions: every phase-1 record is flagged and no record is flagged in phases 2-4. Each gives {integer(gb["tp"])} detections, {integer(gb["fp"])} false alerts and {integer(gb["alerts"])} alerts; XGBoost gives {integer(xb["tp"])}, {integer(xb["fp"])} and {integer(xb["alerts"])} respectively. '
         f'The lower assumed loss therefore reflects {fewer_fp} fewer false alerts at the cost of {fewer_tp} fewer detection. These are three comparisons with the same final rule, not three distinct empirical confirmations. Lower loss than that rule does not establish lower loss than no alarms. These exploratory, unadjusted results are retained without selecting a new preferred workflow.')

    heading('5.3 Historical selection explains capacity and no-alarm outcomes', page=True)
    rows=[]
    for f in range(4):
        a=[r for r in va if match(r, fold=f)]; c=[r for r in vc if match(r, fold=f)]
        quiet=[r for r in silent if match(r, fold=f)]
        rows.append([f+1, f'{sum(float(r["excess"])>0 for r in a)}/4', f'{sum(float(r["excess"])>0 for r in c)}/16',
                     sum(r['no_alarm_reason']=='cost_selects_no_alarm' for r in quiet),
                     sum(r['no_alarm_reason']=='capacity_changes_to_no_alarm' for r in quiet)])
    table('capacity', 'Fixed-XGBoost capacity excess and historical reasons for C no-alarm selection.',
          ['Phase','A excess settings','C excess settings','Cost selects none','Capacity changes to none'], rows)
    para(f'All reference areas contain positive labels, yet {len(silent)} of the 64 C settings select no alarms. Of these, {causes["cost_selects_no_alarm"]} follow unconstrained cost optima and {causes["capacity_changes_to_no_alarm"]} follow changes induced by the historical capacity restriction. '
         'These classifications use reference decisions, not test performance. A small reference advantage over another feasible alarm set describes a weak selection margin; it does not by itself establish unstable thresholds.')
    example=one(xv, fold=1, mechanism='A', budget=.2, cost=10)
    para(f'The budget rule exceeds later capacity in {over_a}/16 phase-budget settings and C in {over_c}/64 phase-budget-cost settings. The settings are dependent and use different denominators. '
         f'In phase 2, A at a 20% budget produces {integer(example["alerts"])} alerts, with {integer(example["excess"])} above its {integer(example["slots"])} later slots. The reference budget had been feasible. '
         'Different reference/test prevalences, including only one positive among 207 reference records in phase 2, provide context but do not identify prevalence shift as the sole cause; class-conditional scores can also change. Batch Top-k enforces capacity within the complete observed score batch and remains a retrospective ranking reference, not a real-time rule or performance bound.')
    para('On the reference area, C cannot have lower optimized loss than B because its candidate set is a subset. That ordering need not survive transfer. Whole-tie rules and fixed numerical cutoffs explain how historical feasibility can coexist with both later underuse and excess demand.')

    heading('5.4 Model updating at a frozen numerical cutoff', page=True)
    rows=[]
    for mechanism in ('A','B','C'):
        q=[r for r in transfers if match(r, model='XGBoost', scheme='time', mechanism=mechanism)]
        decisions={(r['fold'],r['rule']):r for r in all_audits if match(r, model='XGBoost', scheme='time', mechanism=mechanism)}
        finite=[r for r in q if r['threshold_kind']=='finite']
        rows.append([mechanism,sum(math.isfinite(float(r['threshold'])) for r in decisions.values()),
                     sum(float(r['threshold'])==math.inf for r in decisions.values()),
                     sum(float(r['threshold'])==-math.inf for r in decisions.values()),
                     span(r['loss100_delta'] for r in finite), span((r['alerts_delta'] for r in finite), places=0)])
    table('refitting', 'XGBoost refitted-minus-fixed changes. Finite groups are defined by historical selection.',
          ['Rule','Finite','No alarm','All alarm','Finite Delta L100','Finite Delta alerts'], rows)
    para('Distinct historical settings count phase times budget for A, phase times cost for B, and phase times budget times cost for C. A is reused when evaluating costs. The ranges retain all stated evaluation costs. Refitting can alter both the fitted function and raw-score scale; using the same cutoff measures their combined workflow difference, not a causal effect of either component. Infinite cutoffs are invariant by construction and provide no evidence of finite-cutoff robustness.')
    refit=one(transfers,model='XGBoost',scheme='time',fold=0,rule='A_b20',cost=10)
    refit_ci=one(comparisons,model='XGBoost',comparison='Refitted-Fixed',left_rule='A_b20',cost=10,fold=0)
    para(f'For a specific finite cutoff, phase 1 A at budget 20% and r=10 uses {num(refit["threshold_fixed"],6)} in both workflows. Alerts change from {integer(refit["alerts_fixed"])} fixed to {integer(refit["alerts_refitted"])} refitted; detections change from {integer(refit["tp_fixed"])} to {integer(refit["tp_refitted"])}. '
         f'Refitted-minus-fixed loss is {num(refit["loss100_delta"],2)} per 100 records, with paired 95% interval {interval(refit_ci["ci_low"],refit_ci["ci_high"],2)}. This interval concerns model updating, separately from the C-A policy intervals in Section 5.')
    para('The retained block intervals are descriptive and unadjusted across settings. They characterize uncertainty in differences among frozen decisions on the observed queue of records. LR and bagged CART supply additional model comparisons under the same reference boundaries; they do not certify robustness across mines, future distributions or untested model classes. Complete per-phase results and all costs remain in the Stage 2 technical report and saved files.')

    heading('5.5 Frozen-rule cost and prevalence descriptions', page=True)
    rows=[]
    for b in plan['budgets']:
        r=one(boundary, model='XGBoost', scheme='time', cohort='pooled', workflow='fixed', mechanism='C', selection_cost=10, budget=b)
        rows.append([f'{100*b:.0f}%',integer(r['tp']),integer(r['fp']),integer(r['alerts']),
                     '-' if not r['break_even_cost'] else num(r['break_even_cost'],2)])
    table('boundaries', 'Cost equality for unchanged pooled C rules selected historically at rselect=10.',
          ['Historical budget','TP','FP','Alerts','Equality cost r*'], rows)
    para('For a frozen alert set, Delta Lnone,100(reval) = 100(FP - reval TP)/N. When TP>0, equality occurs at r*=FP/TP and the assumed loss favors the unchanged rule above that evaluation cost. TP=FP=0 is equivalent to no alarms for every cost; with TP=0 and FP>0 there is no finite improvement. The dash identifies no-alarm equivalence. These are descriptive equality points on this cohort, not cost estimates or test-selected recommendations.')
    rows=[]
    for b in plan['budgets']:
        cells=[f'{100*b:.0f}%']
        for pi in (.05,.10):
            r=one(scenarios, model='XGBoost', scheme='time', fold=-1, workflow='fixed', cost=10, budget=b,
                  rule=f'C_b{round(100*b):02d}_r10', scenario_prevalence=pi)
            cells += [num(r['delta_loss_vs_no_alarm_100'],2), f'{100*float(r["expected_alert_rate"]):.2f}%',
                      'NA' if not r['expected_precision'] else f'{100*float(r["expected_precision"]):.2f}%']
        rows.append(cells)
    table('scenarios', 'Frozen prevalence scenarios for C selected at r=10. Loss differences are versus no alarms; alerts and precision are expectations.',
          ['Budget','pi 5% Delta','Alerts','Precision','pi 10% Delta','Alerts','Precision'], rows)
    para('The full supplement retains prevalences 2%, 5%, 10% and 15%. The compact display pairs loss and workload at 5% and 10% without reselecting policies. Assuming unchanged within-class rates, L100(pi)=100[r pi FNR+(1-pi)FPR], alert fraction=pi TPR+(1-pi)FPR and precision=pi TPR/alert fraction. Precision is undefined at zero expected alerts (NA); observed zero-alert precision retains its existing zero convention. These scenarios neither diagnose actual shift nor estimate the probability of exceeding capacity.')

    heading('5.6 Monitoring information at matched grades and slots')
    para('Historical threshold transfer changes both detection and workload. This supplement holds workload and grade allocation fixed to ask how much additional hazardous-shift coverage the monitoring variables supply. The four-input baseline contains seismic, seismoacoustic, ghazard and shift type; the extended design uses the retained 17 columns. Both use XGBoost, the same four historical fitting/test splits and the same inner AP search, tuned separately for each design. Neither policy-reference labels nor test labels enter fitting or tuning. The two designs therefore receive matched tuning opportunities, unlike the earlier ratings-only ablation that omits shift type.')
    para('Within each test block, K=floor(B*Ntest) slots are allocated to higher seismic grades first and then to higher model scores within each grade. Both models receive identical total and grade-specific slots. A boundary tie is allocated uniformly: h slots among m tied records containing p hazardous labels contributes E(TP)=h*p/m; labels evaluate the allocation and do not determine it. Counts with decimals are expectations over tie allocation. Each of the 2,000 paired block-bootstrap draws reranks its sampled multiset and recomputes the slots and ties. Length 32 is primary; lengths 16 and 64 and the complete table remain in the supplement. These conditional intervals freeze model fits and tuning and do not cover their uncertainty.')
    para(f'At 20% test-batch slots, the pooled baseline has E(TP)={num(information20["base_etp"],2)} and the extended design {num(information20["extended_etp"],2)}, each using {integer(information20["k"])} slots. The difference is {num(information20["delta_etp"],2)} with conditional 95% interval {interval(information20_ci["ci_low"],information20_ci["ci_high"],2)}. All 16 block-specific primary intervals and all four pooled intervals include zero; the point differences vary across blocks. This workflow has not established a consistent monitoring-information increment, and it does not establish that monitoring information is generally uninformative.')
    figure('reports/overall/monitoring_increment.png','Figure 5. Extended minus grade-plus-shift expected detections at identical seismic-grade allocations and test-batch slots. Positive differences favor added monitoring; bars are conditional 95% paired reranking intervals at block length 32. All four blocks and four slot fractions are shown; pooled results are supplementary. These fractions fix test-batch slots, whereas the earlier historical budgets constrain reference threshold selection and permit future workload to change.',width=6.2)
    para(f'The two uses of 20% answer different questions. Historical C at r=10 produces {integer(c20["alerts"])} test alerts and {integer(c20["tp"])} detections, whereas this batch study allocates {integer(information20["k"])} slots to each model. Its {num(information20["extended_etp"],2)} expected detections cannot be compared with C as an improvement at equal workload. The batch task is retrospective ranking, not observed dispatch. Its protocol was specified after these test records had been examined; the records overlap the two stages and add no independent validation cohort.')
    items.append(dict(kind='link',text='Exploratory supplement: matched design, all intervals and reproduction details',target='../../reports/information/monitoring_information_increment.pdf'))
    items.append(dict(kind='link',text='Supplementary decision combinations and threshold evidence',target='../../reports/overall/comparator_appendix.md'))

    heading('6 Integrated interpretation and next evidence', page=True)
    para('Common-cohort comparisons establish protocol-gap sensitivity, probability correction concerns a named error reference, and feature attribution concerns fitted-model use. These supporting analyses distinguish the forecast evidence needed to interpret subsequent warning decisions.')
    para('Decision evaluation adds information that ranking, probability assessment and attribution do not supply. At the displayed cost, C can lower loss and inspection demand relative to A while having no favorable pooled point difference against no alarms. Historical capacity choices explain why rules are silent or active, whereas later counts show their transferred coverage and workload. Refitting changes which score function is subjected to a historical cutoff. These comparisons establish workflow-specific tradeoffs; they do not identify a universally preferred policy.')
    para('Phase-specific evidence must accompany the pooled comparison. In phase 1, C at 5% and 10% has no detections and loss-difference intervals entirely above zero. The nonzero pooled intervals cross zero. Pooling therefore does not describe every phase, and an unresolved pooled direction does not erase a resolved increase within the stated phase and conditional analysis.')
    para('The shared cohort is also the main inference boundary. Recorded order supplies no verified timestamps or longwall continuity. Deduplication reproduces a public mirror, not a certified operational history. Costs, capacity slots and the absence of an excess-demand penalty are explicit scenario assumptions. The target counts hazardous shifts, not accidents, executed inspections or interventions. Two analyses of the same observed records increase explanatory coverage, not the number of independent validation cohorts.')
    para('Additional direct-grade and no-alarm contrasts show that decision value depends on the comparator: cost-only XGBoost has lower loss than the identical grade-rule decisions at r=10, while its no-alarm interval spans zero. The matched-information analysis fixes workload and does not establish a consistent monitoring increment beyond grades and shift type. Full refit/reselection uncertainty requires a separate protocol; credible lag features require verified continuity. Independent-site and inspection outcomes are the next evidence needed.')
    heading('7 Conclusions')
    para(f'Comparable-cohort evaluation reduces the original AP gaps by {reductions}; this quantifies cohort sensitivity, not leakage or drift attribution. At r=10 and a 20% historical budget, capacity-constrained cost rules produce {integer(c20["alerts"])} alerts and {integer(c20["tp"])}/{positives} detections, versus {integer(a20["alerts"])} alerts and {integer(a20["tp"])}/{positives} for budget-only rules. Phase 1 C at 5% and 10% has no detections and positive loss-difference intervals; the nonzero pooled C intervals span zero. Joint loss, detection and workload accounting shows why historical optimality and capacity feasibility do not establish later decision value.')
    para(f'Fixing grade allocations and test-batch slots provides a different comparison: at 20%, monitoring inputs add {num(information20["delta_etp"],2)} expected detections to a grade-plus-shift model, with conditional interval {interval(information20_ci["ci_low"],information20_ci["ci_high"],2)} and {integer(information20["k"])} slots per model. The block-specific evidence does not establish a consistent increment. Historical threshold transfer and equal-slot batch ranking should therefore remain distinct when interpreting decision value. Independent timestamped working-face data with inspection outcomes are needed to test transfer beyond this cohort.')
    heading('Data, code, contributions and AI assistance')
    para('The repository retains source mapping, predictions, historical candidates, row-level decisions, resampling outputs and version-specific checks. The reporting guide documents generation and verification; both stage papers provide technical detail. Code and evidence: https://github.com/junjie-yang11/seismic-bump-forecasting.')
    para('The author directed scope, protocol refinement and review priorities. ChatGPT and Codex substantially assisted implementation, execution, computational and reference checks, and writing. CONTRIBUTIONS.md documents these roles. AI-assisted writing workflows do not supply scientific evidence for the research conclusions.')
    heading('References', page=True)
    for index,key in enumerate(used,1):
        text=references[key].replace('doi:','DOI: ').replace('; 2023;108:', ', 2023, 108:').replace('; 1989;17(3):', ', 1989, 17(3):').replace('1217?1241','1217-1241')
        para(f'[{index}] {text}')
    return items, tables, evidence.inputs, certificates


def write_outputs(items, destination=None, stem=STEM, title=TITLE):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from phase1.generate_report import scrub_report_package
    destination=Path(destination) if destination is not None else DEST
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.75)
    sec.left_margin=sec.right_margin=Inches(.8)
    normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(11.5)
    normal.paragraph_format.line_spacing=1.12; normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for name,size in [('Title',16),('Heading 1',12.5),('Heading 2',12),('Caption',10)]:
        style=doc.styles[name]; style.font.name='Times New Roman'; style.font.size=Pt(size)
        style.font.color.rgb=None
        if name!='Caption':style.font.bold=True
    doc.styles['Caption'].font.bold=False
    doc.styles['Caption'].font.italic=False
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        if style.type==1:
            fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
            for key in list(fonts.attrib):
                if key.endswith('Theme'):del fonts.attrib[key]
            for key in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+key),'Times New Roman')
    settings=doc.settings.element
    settings.append(OxmlElement('w:doNotAutoCompressPictures'))
    dpi=OxmlElement('w14:defaultImageDpi');dpi.set(qn('w14:val'),'330');settings.append(dpi)
    md=[];break_before=False
    for item in items:
        kind=item['kind']
        if kind=='break':
            break_before=True; md.append('<!-- page break -->'); continue
        if kind=='title':
            p=doc.add_paragraph(item['text'],'Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            md.append('# '+item['text']); continue
        if kind=='heading':
            p=doc.add_paragraph(item['text'],'Heading 1')
            if break_before:p.paragraph_format.page_break_before=True;break_before=False
            md.append('## '+item['text']); continue
        if kind=='equation':
            from reporting.word_math import append_equation, markdown_equation
            if doc.paragraphs:doc.paragraphs[-1].paragraph_format.keep_with_next=True
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_together=True
            append_equation(p,item['id']);md.append(markdown_equation(item['id']));continue
        if kind=='link':
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            from lxml import etree
            p=doc.add_paragraph();link=OxmlElement('w:hyperlink')
            url='https://github.com/junjie-yang11/seismic-bump-forecasting/blob/main/'+item['target'].replace('../../','')
            link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');text=OxmlElement('w:t');text.text=item['text'];run.append(text);link.append(run);p._p.append(link)
            md.append('['+item['text']+']('+item['target']+')');continue
        if kind=='paragraph':
            text=item['text']; p=doc.add_paragraph(text)
            if text in ('Junjie Yang','Mining Engineering, Fuzhou University'):
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            if re.match(r'^\[\d+\]',text):
                p.paragraph_format.left_indent=Inches(.23); p.paragraph_format.first_line_indent=Inches(-.23)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:run.font.size=Pt(10)
            md.append(text); continue
        if kind=='figure':
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            inline=p.add_run().add_picture(str(ROOT/item['path']),width=Inches(item['width']))
            vector=(ROOT/item['path']).with_suffix('.svg')
            if item['path']=='reports/overall/evidence_overview.png' and vector.exists():
                from reporting.visuals import add_svg_alternative
                add_svg_alternative(inline,vector,p.part)
            caption=doc.add_paragraph(item['caption'],'Caption')
            caption.paragraph_format.keep_together=True
            target=Path('../../')/item['path']
            md.append(f'![{item["caption"]}]({target.as_posix()})'); continue
        p=doc.add_paragraph(item['caption'],'Caption'); p.paragraph_format.keep_with_next=True
        headers=item['headers']; rows=item['rows']
        md.append(item['caption'])
        md_rows=['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']
        t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style='Table Grid'
        for c,h in zip(t.rows[0].cells,headers):c.text=h
        for values in rows:
            for c,v in zip(t.add_row().cells,values):c.text=v
            md_rows.append('| '+' | '.join(values)+' |')
        md.append('\n'.join(md_rows))
        for ri,row in enumerate(t.rows):
            prop=row._tr.get_or_add_trPr();prop.append(OxmlElement('w:cantSplit'))
            if ri==0:prop.append(OxmlElement('w:tblHeader'))
            for ci,cell in enumerate(row.cells):
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if ri==0:
                    shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'EDEFF1');cell._tc.get_or_add_tcPr().append(shading)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.02
                    p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.keep_with_next=(ri<len(t.rows)-1 if len(rows)<=10 or item['id'] in
                        ('grade_decisions','ranking_primary','ranking_aux') else ri<2)
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 else (WD_ALIGN_PARAGRAPH.LEFT if ci==0 or item['id']=='protocols' else WD_ALIGN_PARAGRAPH.RIGHT)
                    for run in p.runs:run.font.size=Pt(9.3 if len(headers)>6 else 10);run.bold=ri==0
        widths={'no_alarm_intervals':[.6,1.3,2.3,.6,.6,.9],
                'grade_decisions':[1.0,1.0,1.7,.4,.4,.4,.55,1.05],
                'grade_intervals':[.7,1.2,1.5,1.5,1.5],
                'ranking_primary':[.55,.65,.45,1.,1.,2.75],
                'ranking_aux':[.6,.5,1.,1.,3.3]}.get(item['id'])
        if widths:
            t.autofit=False
            for column,width in zip(t.columns,widths):column.width=Inches(width)
            for row in t.rows:
                for cell,width in zip(row.cells,widths):cell.width=Inches(width)
        doc.add_paragraph().paragraph_format.space_after=Pt(1)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title=title
    doc.save(destination/(stem+'.docx'));scrub_report_package(destination/(stem+'.docx'))
    (destination/(stem+'.md')).write_bytes(('\n\n'.join(md)+'\n').encode('utf-8'))


def write_comparison_appendix():
    """Keep all planned combinations outside the compact main manuscript."""
    def read(name):
        return list(csv.DictReader((ROOT/'results/phase2/comparators'/name).open(encoding='utf-8',newline='')))
    evaluations=read('evaluations.csv');intervals=read('intervals.csv')
    lookup={(r['left_method'],r['rule'],r['cost'],r['fold']):r for r in intervals if r['comparison']=='policy-none'}
    lines=['# Frozen decision comparison appendix',
           'This generated appendix retains all planned A/B/C settings. Phase results precede pooled summaries. Negative loss differences favor the left policy. Intervals are unadjusted 95% percentile intervals from 2,000 paired within-phase moving-block draws of length 32, conditional on frozen policies and these records. They exclude refitting and selection uncertainty. Exact zero intervals can follow identical decisions and do not establish transfer stability.',
           '[Protocol](../../phase2/comparators/protocol.json) · [Evidence and verification](../../results/phase2/comparators/README.md) · [Main manuscript](integrated_research_report.md)',
           '## Policy minus no alarms']
    for cost in (5,10,20,50):
        lines.append('### Assumed missed-shift cost '+str(cost))
        table=['| Phase | Method | Rule | Delta loss / 100 [95%] | TP | FN | FP | Alerts | Excess |',
               '| --- | --- | --- | --- | --- | --- | --- | --- | --- |']
        part=[r for r in evaluations if float(r['cost'])==cost]
        for r in sorted(part,key=lambda r:(int(r['fold']) if int(r['fold'])>=0 else 4,r['method'],r['rule'])):
            ci=lookup[(r['method'],r['rule'],r['cost'],r['fold'])]
            cells=['Pooled' if int(r['fold'])<0 else str(int(r['fold'])+1),r['method'],r['rule'],
                   num(ci['loss_delta100'],2)+' '+interval(ci['ci_low'],ci['ci_high'],2)]
            cells += [integer(r[k]) for k in ('tp','fn','fp','alerts')]
            cells += [integer(r['excess']) if r['excess'] else '-']
            table.append('| '+' | '.join(cells)+' |')
        lines.append('\n'.join(table))
    lines.append('## Fixed XGBoost minus each grade rule')
    for cost in (5,10,20,50):
        lines.append('### Assumed missed-shift cost '+str(cost))
        table=['| Phase | Grade field | Rule | ML minus grade loss / 100 [95%] |',
               '| --- | --- | --- | --- |']
        part=[r for r in intervals if r['comparison']=='ML-grade' and float(r['cost'])==cost]
        for r in sorted(part,key=lambda r:(int(r['fold']) if int(r['fold'])>=0 else 4,r['right_method'],r['rule'])):
            cells=['Pooled' if int(r['fold'])<0 else str(int(r['fold'])+1),r['right_method'],r['rule'],
                   num(r['loss_delta100'],2)+' '+interval(r['ci_low'],r['ci_high'],2)]
            table.append('| '+' | '.join(cells)+' |')
        lines.append('\n'.join(table))
    (DEST/'comparator_appendix.md').write_text('\n\n'.join(lines)+'\n',encoding='utf-8',newline='\n')


def build():
    DEST.mkdir(parents=True,exist_ok=True)
    certificate=DEST/'verification.json'
    started=datetime.now(timezone.utc).isoformat()
    certificate.write_text(json.dumps(dict(passed=False,status='building',started_at_utc=started)),encoding='utf-8')
    try:
        # A prior PDF cannot certify a newly generated Word manuscript.
        (DEST/(STEM+'.pdf')).unlink(missing_ok=True)
        overview_figure(DEST/'evidence_overview.png')
        items,tables,inputs,source_checks=assemble()
        write_outputs(items)
        write_comparison_appendix()
        manifest=dict(title=TITLE,scope='Integrated reporting with frozen-policy, direct-grade and exploratory matched-information evidence; no ML fitting or policy selection in reporting',
                      inputs=inputs,generator_sha256=digest(Path(__file__)),tables=tables,
                      source_certificates_at_build=source_checks,
                      outputs={name:digest(DEST/name) for name in (STEM+'.md',STEM+'.docx','evidence_overview.png','decision_tradeoffs.png','decision_tradeoffs.csv','comparator_appendix.md','evidence_overview.svg','calibration_detail.png','shap_readable.png','monitoring_increment.png','monitoring_increment.csv')})
        (DEST/'report_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
        certificate.write_text(json.dumps(dict(passed=False,status='awaiting independent verification')),encoding='utf-8')
        print(f'Integrated manuscript generated: {len(tables)} tables; model and decision evidence untouched.')
    except (Exception,KeyboardInterrupt) as error:
        certificate.write_text(json.dumps(dict(passed=False,status='failed',started_at_utc=started,
                                               error_type=type(error).__name__,error=str(error)),indent=2),encoding='utf-8')
        raise


if __name__=='__main__':
    build()
