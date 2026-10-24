"""Build a compact matched-information supplement from verified saved evidence."""
from __future__ import annotations
import csv
import hashlib
import json
from pathlib import Path
from datetime import datetime, timezone
from .overall import write_outputs

ROOT=Path(__file__).resolve().parents[1]
DEST=ROOT/'reports/information'
STEM='monitoring_information_increment'
TITLE='Monitoring Information Beyond Hazard Grades and Shift Type'


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def one(rows,**query):
    selected=[r for r in rows if all(float(r[k])==value if isinstance(value,(float,int)) else r[k]==value for k,value in query.items())]
    if len(selected)!=1:raise ValueError(f'Expected one evidence row: {query}; found {len(selected)}')
    return selected[0]


def num(value):return f'{float(value):.2f}'


def effect(row):return num(row['delta_etp'])+' ['+num(row['ci_low'])+', '+num(row['ci_high'])+']'


def increment_figure(rows,path):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    colors={16:'#b36b13',32:'#24668b',64:'#22816d'}
    styles={16:'--',32:'-',64:':'}
    offsets={16:-.16,32:0,64:.16}
    fig,axes=plt.subplots(5,1,figsize=(7.5,7.5),sharex=True)
    low=min(float(r['ci_low']) for r in rows);high=max(float(r['ci_high']) for r in rows)
    margin=(high-low)*.08
    for index,fold in enumerate((0,1,2,3,-1)):
        ax=axes[index]
        title='Pooled (supplementary)' if fold<0 else f'Record block {fold+1}'
        ax.set_title(title,loc='left',fontsize=10,pad=3)
        for j,budget in enumerate((.01,.05,.10,.20)):
            for length in (16,32,64):
                r=one(rows,fold=fold,budget=budget,block_length=length)
                ax.plot([float(r['ci_low']),float(r['ci_high'])],[j+offsets[length]]*2,
                        color=colors[length],ls=styles[length],lw=1.8 if length==32 else 1.3)
            point=one(rows,fold=fold,budget=budget,block_length=32)
            ax.plot(float(point['delta_etp']),j,'o',ms=4,color='#192a35')
        ax.axvline(0,color='#7b858e',ls='--',lw=.8)
        ax.set_xlim(low-margin,high+margin);ax.set_ylim(3.55,-.55)
        ax.set_yticks(range(4),['1%','5%','10%','20%']);ax.set_ylabel('Budget',fontsize=9)
        ax.spines[['top','right']].set_visible(False)
        ax.tick_params(labelsize=9);ax.grid(axis='x',alpha=.13);ax.set_axisbelow(True)
        if fold<0:ax.set_facecolor('#f0f3f5')
    legend=[Line2D([0],[0],color=colors[length],ls=styles[length],lw=1.8 if length==32 else 1.3,
                   label=f'L={length}'+(' (main)' if length==32 else '')) for length in (16,32,64)]
    fig.legend(handles=legend,loc='upper right',bbox_to_anchor=(.99,1.005),frameon=False,ncol=3,fontsize=9)
    axes[-1].set_xlabel('Extended minus base expected hazardous-shift detections',fontsize=10)
    fig.tight_layout(rect=(0,0,1,.975),h_pad=.65)
    fig.savefig(path,dpi=240,bbox_inches='tight');plt.close(fig)


def format_table(destination):
    """Set widths on the new document only; leave the shared writer unchanged."""
    from docx import Document
    from docx.shared import Inches,Pt
    from phase1.generate_report import scrub_report_package
    path=destination/(STEM+'.docx');doc=Document(path)
    table=doc.tables[0];table.autofit=False
    widths=[.7,.65,.35,1.25,1.25,2.35]
    for column,width in zip(table.columns,widths):column.width=Inches(width)
    for row_index,row in enumerate(table.rows):
        for cell,width in zip(row.cells,widths):
            cell.width=Inches(width)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before=paragraph.paragraph_format.space_after=Pt(2.4)
                # Keep the supplementary four-budget pooled group together.
                # Phase results may continue normally with repeated headers.
                if row_index>=17:
                    paragraph.paragraph_format.keep_with_next=row_index<20
                for run in paragraph.runs:run.font.size=Pt(9.5)
    doc.save(path);scrub_report_package(path)


def build(root=ROOT):
    root=Path(root);dest=root/'reports/information';dest.mkdir(parents=True,exist_ok=True)
    record=dest/'verification.json';started=datetime.now(timezone.utc).isoformat()
    record.write_text(json.dumps(dict(passed=False,status='building',started_at_utc=started)),encoding='utf-8')
    (dest/(STEM+'.pdf')).unlink(missing_ok=True)
    try:
        inputs={}
        def read(relative):
            inputs[relative]=sha(root/relative)
            return (root/relative).read_text(encoding='utf-8')
        def csv_input(name):return list(csv.DictReader(read('results/phase2/information/'+name+'.csv').splitlines()))
        plan=json.loads(read('phase2/information/protocol.json'))
        verification=json.loads(read('results/phase2/information/verification.json'))
        if not verification.get('passed') or not verification.get('refit'):
            raise RuntimeError('Passed separate numerical verification including full refitting is required')
        read('results/phase2/information/run_manifest.json')
        for name in ('phase2/information/run.py','phase2/information/verify.py'):
            read(name)
        ev=csv_input('evaluations');intervals=csv_input('intervals');csv_input('rows')
        if len(ev)!=20 or len(intervals)!=60:raise ValueError('Complete evaluation and block-length grid required')
        main_length=plan['bootstrap']['main_block_length']
        figure_rows=[]
        for fold in (0,1,2,3,-1):
            for budget in plan['budgets']:
                for length in plan['bootstrap']['block_lengths']:
                    r=one(intervals,fold=fold,budget=budget,block_length=length)
                    figure_rows.append({key:r[key] for key in ('fold','budget','block_length','replicates','delta_etp','ci_low','ci_high')})
        with (dest/'figure_data.csv').open('w',encoding='utf-8',newline='') as stream:
            writer=csv.DictWriter(stream,fieldnames=list(figure_rows[0]));writer.writeheader();writer.writerows(figure_rows)
        increment_figure(figure_rows,dest/'information_increment.png')
        rows=[]
        for fold in (0,1,2,3,-1):
            for budget in plan['budgets']:
                r=one(ev,fold=fold,budget=budget);ci=one(intervals,fold=fold,budget=budget,block_length=main_length)
                rows.append(['Pooled' if fold<0 else str(fold+1),f'{100*budget:.0f}%',str(int(float(r['k']))),
                             num(r['base_etp'])+' / '+num(r['base_efn']),num(r['extended_etp'])+' / '+num(r['extended_efn']),effect(ci)])
        table=dict(kind='table',id='information_primary',caption='Table 1. Matched information designs at identical total and seismic-grade allocations. TP and FN are analytic expectations; K is the exact alert count. Positive differences favor the extended design. Brackets are conditional 95% intervals at block length 32.',
                   headers=['Block','Budget','K','Base E(TP) / E(FN)','Extended E(TP) / E(FN)','Delta E(TP) [95% CI]'],rows=rows)
        items=[dict(kind='title',text=TITLE),dict(kind='paragraph',text='Junjie Yang')]
        def heading(text):items.append(dict(kind='heading',text=text))
        def para(text):items.append(dict(kind='paragraph',text=text))
        pooled=one(ev,fold=-1,budget=.2);pooled_ci=one(intervals,fold=-1,budget=.2,block_length=main_length)
        all_pooled_include_zero=all(float(r['ci_low'])<=0<=float(r['ci_high']) for r in intervals if int(r['fold'])==-1)
        heading('Question and matched comparison')
        para('Does monitoring information improve hazardous-shift detection beyond three existing hazard assessments and shift type, at identical inspection allocations? This separately specified exploratory supplement compares a four-input base design with the established 17-column design. Both use XGBoost and the same four historical fitting/test splits. Each design independently selects depth 2/3 and 80/160 rounds on the same inner historical split by AP, with negative Brier if that inner set has no positives. Policy-reference and test labels do not enter fitting or tuning; no holdout is added.')
        para('Within each test block, K=floor(B*N) slots go to higher seismic grades first; model scores rank records within each grade. Both models use uniform allocation within an exactly tied boundary score group: h slots among m records containing p hazards contribute E(TP)=h*p/m. Reported fractional counts integrate allocation randomness rather than describe fractional shifts. Pooled results sum block allocations; scores are never ranked across fitted blocks.')
        heading('Expected detections on the shared test records')
        para(f'The four blocks contain 516, 515, 516 and 516 records and 36, 10, 24 and 18 hazardous labels, respectively: 2,063 records and 88 labels in total. At the 20% budget, the pooled extended design detects {num(pooled["extended_etp"])} versus base {num(pooled["base_etp"])} expected hazards with {int(float(pooled["k"]))} alerts each; the difference is {effect(pooled_ci)}. '+('All pooled intervals across the three specified block lengths include zero.' if all_pooled_include_zero else 'Pooled interval directions are retained separately for every specified block length.'))
        items.append(table)
        heading('Block-length sensitivity')
        para('The saved scores, tuning choices and models are frozen. Each of 2,000 paired moving-block draws reconstructs a queue separately within each block, recomputes grade allocations and reranks both score sets, including analytic tie expectations. The same sampled records feed both designs and every budget. Zero-positive draws are retained. Length 32 is primary; lengths 16 and 64 are displayed without selecting one from the outcomes. These are record-order blocks, not verified physical time durations. Intervals exclude fitting, tuning and selection uncertainty.')
        heading('Discussion')
        positive_counts=[]
        for fold in range(4):
            positive_counts.append(str(sum(float(r['delta_etp'])>0 for r in ev if int(r['fold'])==fold))+'/4')
        main_intervals=[r for r in intervals if int(r['fold'])>=0 and int(r['block_length'])==main_length]
        include_count=sum(float(r['ci_low'])<=0<=float(r['ci_high']) for r in main_intervals)
        para('The matched baseline changes the interpretation of the earlier complete-score-versus-grade-random pilot: the present increment is specifically relative to a fitted model already using all three assessments and shift type. Positive point differences occur in '+', '.join(positive_counts)+' budgets in blocks 1-4, respectively; '+str(include_count)+'/'+str(len(main_intervals))+' main block-specific intervals include zero. '+('The pooled intervals also include zero at every budget and block length, so these observations do not resolve a consistent monitoring-information increment under the specified workflow. ' if all_pooled_include_zero else 'The displayed interval directions bound the evidence for increments under the specified workflow. ')+
             'This does not establish that monitoring information is generally uninformative: the comparison concerns these inputs, separately AP-selected configurations and previously examined records. At fixed P and K, E(L)=r*P+K-(r+1)*E(TP), so an additional cost grid would restate detection differences. Recorded order remains a temporal proxy, and the batch allocation is not an observed inspection process or accident-prevention effect. Independent timestamped working-face data with inspection outcomes would test transfer beyond this cohort.')
        heading('Protocol and computational evidence')
        para('The separate protocol is phase2/information/protocol.json. Model fitting and analytic allocation are implemented in phase2/information/run.py; phase2/information/verify.py recomputes the evidence through a separate calculation path and supports full refitting replay. Results in results/phase2/information/ retain row-level scores and inclusion probabilities, all candidate tuning predictions, eight final models, grade allocations, every bootstrap multiplicity and replicate, and version-specific verification. The original two-stage papers and ranking pilot remain unchanged. Author contributions and AI assistance follow CONTRIBUTIONS.md.')
        heading('Increment and interval figure')
        items.append(dict(kind='figure',path='reports/information/information_increment.png',width=6.4,
            caption='Figure 1. Extended minus base expected detections at identical allocations. Dots are the unchanged original-queue point differences. Lines show paired 95% intervals for lengths 16 (orange dashed), 32 (blue solid, main) and 64 (green dotted). Positive favors added monitoring. The shaded pooled panel is supplementary; block-specific evidence is primary.'))
        write_outputs(items,dest,STEM,TITLE)
        format_table(dest)
        manifest=dict(title=TITLE,created_at_utc=datetime.now(timezone.utc).isoformat(),inputs=inputs,
            generator_sha256=sha(Path(__file__)),writer_sha256=sha(root/'reporting/overall.py'),tables=[{k:v for k,v in table.items() if k!='kind'}],
            scope='Separate matched-information report from computed evidence; one phase table, one all-length interval figure, no fitting or policy selection by reporting',
            numerical_verification_at_build=dict(sha256=sha(root/'results/phase2/information/verification.json'),refit=verification['refit'],scope=verification['scope']),
            outputs={name:sha(dest/name) for name in (STEM+'.md',STEM+'.docx','figure_data.csv','information_increment.png')})
        (dest/'report_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
        record.write_text(json.dumps(dict(passed=False,status='awaiting separate report verification')),encoding='utf-8')
        print('Information supplement generated: one table, one figure, original papers unchanged.')
    except (Exception,KeyboardInterrupt) as error:
        record.write_text(json.dumps(dict(passed=False,status='failed',started_at_utc=started,error_type=type(error).__name__,error=str(error)),indent=2),encoding='utf-8');raise


if __name__=='__main__':build()
