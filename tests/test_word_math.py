"""Editable equation structure is evidence, not just flattened equation text."""
from copy import deepcopy
from io import BytesIO
import unittest
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from reporting import word_math as math


class WordMathTests(unittest.TestCase):
    def test_equations_round_trip_structurally(self):
        from docx import Document
        for name, equation in math.EQUATIONS.items():
            with self.subTest(name=name):
                node = math.append_equation(Document().add_paragraph(), name)
                self.assertEqual(math.math_signature(node), equation['ast'])
                # Exercise the standard-library XML reader used by verifiers.
                self.assertEqual(math.math_signature(ET.fromstring(ET.tostring(node))), equation['ast'])

    def test_native_fractions_subscripts_and_floor_are_structural(self):
        from docx import Document
        ns = {'m': math.MATH_NS}
        loss = math.append_equation(Document().add_paragraph(), 'loss100')
        self.assertEqual(len(loss.findall('.//m:f', ns)), 1)
        self.assertEqual(loss.find('.//m:f/m:den/m:r/m:t', ns).text, 'N')
        self.assertEqual(loss.find('.//m:sSub/m:sub/m:r/m:t', ns).text, '100')
        excess = math.append_equation(Document().add_paragraph(), 'excess')
        floors = excess.findall('.//m:begChr', ns)
        self.assertIn('⌊', [x.get('{' + math.MATH_NS + '}val') for x in floors])

    def test_changed_denominator_minus_and_subscript_change_signature(self):
        from docx import Document
        ns = {'m': math.MATH_NS}
        original = math.append_equation(Document().add_paragraph(), 'delta_none100')
        for path, replacement in [('.//m:den/m:r/m:t', 'P'),
                                  ('.//m:sSub/m:sub/m:r/m:t', 'none,10')]:
            mutated = deepcopy(original); mutated.find(path, ns).text = replacement
            self.assertNotEqual(math.math_signature(mutated), math.EQUATIONS['delta_none100']['ast'])
        mutated = deepcopy(original)
        minus = next(x for x in mutated.findall('.//m:t', ns) if x.text == '−')
        minus.text = '+'
        self.assertNotEqual(math.math_signature(mutated), math.EQUATIONS['delta_none100']['ast'])

    def test_unknown_id_and_unrecognized_operators_fail(self):
        from docx import Document
        with self.assertRaises(KeyError):
            math.append_equation(Document().add_paragraph(), 'unknown')
        with self.assertRaises(KeyError):
            math.markdown_equation('unknown')
        node = ET.Element('{' + math.MATH_NS + '}oMath')
        ET.SubElement(node, '{' + math.MATH_NS + '}rad')
        with self.assertRaises(ValueError):
            math.math_signature(node)

    def test_saved_docx_preserves_editable_math_and_markdown(self):
        from docx import Document
        document = Document()
        for name in math.EQUATIONS:
            math.append_equation(document.add_paragraph(), name)
            self.assertTrue(math.markdown_equation(name).startswith('$$\n'))
        package_bytes = BytesIO(); document.save(package_bytes)
        with ZipFile(package_bytes) as package:
            tree = ET.fromstring(package.read('word/document.xml'))
            self.assertFalse(any(name.startswith('word/media/') for name in package.namelist()))
        nodes = tree.findall('.//{' + math.MATH_NS + '}oMath')
        self.assertEqual([math.math_signature(x) for x in nodes],
                         [x['ast'] for x in math.EQUATIONS.values()])


if __name__ == '__main__':
    unittest.main()
