"""Generate the report from result CSVs, with one source for Markdown and Word.

Run experiments first. Word creation needs requirements-report.txt. On Windows,
export_report.ps1 exports the Word report to PDF using installed Microsoft Word.
"""
from __future__ import annotations
import csv
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

def scrub_report_package(path):
    """Remove unused template stores and revision identifiers from the report."""
    from zipfile import ZipFile, ZIP_DEFLATED
    from lxml import etree
    temporary = path.with_suffix('.clean.docx')
    with ZipFile(path) as source, ZipFile(temporary, 'w', ZIP_DEFLATED) as target:
        for name in source.namelist():
            if name.startswith('customXml/') or name == 'docProps/custom.xml':
                continue
            payload = source.read(name)
            if name.endswith(('.xml', '.rels')):
                tree = etree.fromstring(payload)
                for element in list(tree.iter()):
                    local = etree.QName(element).localname
                    if local == 'Relationship' and 'customXml' in element.get('Type', ''):
                        element.getparent().remove(element)
                    elif local == 'Override' and (element.get('PartName', '').startswith('/customXml/')
                                                 or element.get('PartName') == '/docProps/custom.xml'):
                        element.getparent().remove(element)
                    elif local == 'rsids' and element.getparent() is not None:
                        element.getparent().remove(element)
                    else:
                        for key in list(element.attrib):
                            if etree.QName(key).localname.startswith('rsid'):
                                del element.attrib[key]
                payload = etree.tostring(tree, xml_declaration=True, encoding='UTF-8', standalone=True)
            target.writestr(name, payload)
    temporary.replace(path)

def rows(name):
    with (ROOT / 'results' / name).open(encoding='utf-8', newline='') as f:
        return list(csv.DictReader(f))

def short(model):
    return 'LR' if 'logistic' in model else 'CART'

def table(headers, data):
    return '\n'.join(['| ' + ' | '.join(headers) + ' |',
        '| ' + ' | '.join(['---'] * len(headers)) + ' |'] +
        ['| ' + ' | '.join(map(str, row)) + ' |' for row in data])

def number(value):
    return '{:.4f}'.format(float(value))

def report_markdown():
    from scripts.paper_content import build_paper
    return build_paper(ROOT)

def write_markdown_report():
    (ROOT / 'report').mkdir(exist_ok=True)
    (ROOT / 'report/technical_note.md').write_text(report_markdown(), encoding='utf-8')

def write_word_report():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    doc = Document()
    # Reports for external sharing contain project content only.
    for field_name in ('author', 'last_modified_by', 'comments', 'keywords',
                       'subject', 'identifier', 'category', 'content_status'):
        setattr(doc.core_properties, field_name, '')
    privacy = OxmlElement('w:removePersonalInformation')
    doc.settings.element.append(privacy)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.85)
    section.left_margin = section.right_margin = Inches(1)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption'):
        s = doc.styles[name]
        s.font.name = 'Times New Roman'
        s.font.color.rgb = RGBColor(0, 0, 0)
        for attr in ('ascii','hAnsi','eastAsia','cs'):
            s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:'+attr),'Times New Roman')
        fonts=s.element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            if 'theme' in key.lower(): del fonts.attrib[key]
        for node in list(s.element.iter(qn('w:pBdr'))): node.getparent().remove(node)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.widow_control = True
    doc.styles['Title'].font.size = Pt(17)
    doc.styles['Title'].font.bold = True
    doc.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles['Title'].paragraph_format.space_after = Pt(12)
    doc.styles['Heading 1'].font.size = Pt(12)
    doc.styles['Heading 2'].font.size = Pt(11)
    for name in ('Heading 1','Heading 2'):
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.space_before = Pt(12)
        doc.styles[name].paragraph_format.space_after = Pt(5)
        doc.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.styles['Caption'].font.size = Pt(9.5)
    doc.styles['Caption'].font.bold = False
    source = report_markdown()
    lines = source.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line:
            i += 1; continue
        if line.startswith('|'):
            data = []
            while i < len(lines) and lines[i].startswith('|'):
                data.append([v.strip() for v in lines[i].strip('|').split('|')]); i += 1
            data.pop(1)
            tbl = doc.add_table(rows=0, cols=len(data[0]))
            tbl.autofit = False
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            weights = [max(1.,min(2.5,len(h)/12.)) for h in data[0]]
            if data[0][0]=='Model': weights[0]=1.15
            if data[0][-1]=='95% interval': weights[-1]=2.
            if data[0][0]=='Model / variant': weights[0]=2.2
            if data[0][:2]==['Model', 'Feature set']: weights=[.9, 2.4, 1., 1., 1., 1.]
            if data[0]==['Feature', 'Group', 'Definition', 'Encoding']: weights=[1.1, 1.1, 3.4, .9]
            widths=[6.5*w/sum(weights) for w in weights]
            for col,width in zip(tbl.columns,widths): col.width=Inches(width)
            for ri, values in enumerate(data):
                cells = tbl.add_row().cells
                for cell, value, width in zip(cells, values, widths):
                    cell.text = value
                    cell.width = Inches(width)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    pr = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement('w:tcBorders')
                    for edge in ('top', 'left', 'bottom', 'right'):
                        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single')
                        e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
                    pr.append(borders)
                    if ri == 0:
                        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'F2F2F2'); pr.append(shading)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(4)
                        p.paragraph_format.space_after = Pt(4)
                        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.keep_with_next = ri < (2 if len(data) > 10 else len(data)-1)
                        for r in p.runs: r.font.size = Pt(9.5); r.bold = ri == 0
                trpr = tbl.rows[-1]._tr.get_or_add_trPr()
                trpr.append(OxmlElement('w:cantSplit'))
                if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        image = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
        if image:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            figure_width = 5.8 if 'reliability' in image.group(2) else (6.2 if 'shap_' in image.group(2) else 5.3)
            p.add_run().add_picture(str((ROOT / 'report' / image.group(2)).resolve()), width=Inches(figure_width))
            caption = doc.add_paragraph(image.group(1), style='Caption')
            caption.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif line.startswith('# '): doc.add_paragraph(line[2:], style='Title')
        elif line.startswith('## '):
            p=doc.add_paragraph(line[3:],style='Heading 1')
            if line[3:]=='Abstract': p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('### '): doc.add_paragraph(line[4:], style='Heading 2')
        elif re.match(r'^Table (?:\d+|A\d+)\.',line):
            p=doc.add_paragraph(line,style='Caption')
            p.paragraph_format.keep_with_next=True
        else:
            p=doc.add_paragraph(line)
            # Only the byline block preceding Abstract is centered.
            if '## Abstract' in lines and i < lines.index('## Abstract'):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs: r.font.size = Pt(10.5)
            if line.startswith('The scatter distributions expose'):
                p.paragraph_format.keep_together = True
            if line.startswith('Table 5 reports fixed-threshold holdout performance.'):
                p.paragraph_format.keep_together = True
            if line.startswith('Most of the apparent'):
                p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            if line.startswith('Keywords:'):
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs: r.font.size=Pt(10)
            if re.match(r'^\[\d+\]',line):
                p.paragraph_format.left_indent=Inches(.22)
                p.paragraph_format.first_line_indent=Inches(-.22)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs: r.font.size=Pt(10)
        i += 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    output = ROOT / 'report/technical_report.docx'
    doc.save(str(output))
    scrub_report_package(output)

if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser(); ap.add_argument('--docx', action='store_true')
    args = ap.parse_args()
    write_markdown_report()
    if args.docx: write_word_report()
