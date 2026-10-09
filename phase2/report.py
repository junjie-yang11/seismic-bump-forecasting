"""A standalone academic report authored entirely from saved phase-two evidence.

Run with the bundled document runtime and the separate plotting dependencies.
No model is fitted and no experiment or first-stage output is modified here.
"""
import hashlib
import json
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .plan import ROOT, OUT, PLAN

DEST=ROOT/'reports/phase2'

def build():
    DEST.mkdir(parents=True,exist_ok=True)
    read=lambda n:pd.read_csv(OUT/(n+'.csv'))
    e,a,c,p=map(read,['evaluations','audits','comparisons','predictions'])
    verification=json.loads((OUT/'verification.json').read_text())
    if not verification['passed'] or not verification['replay']:
        raise RuntimeError('Independent model replay is required before reporting results')
    if hashlib.sha256((ROOT/'phase2/verify.py').read_bytes()).hexdigest()!=verification['verifier_sha256']:
        raise RuntimeError('Main verifier changed: rerun independent model replay before reporting')
    for name,digest in {**verification['evidence_sha256'],**verification['provenance_sha256']}.items():
        if hashlib.sha256((OUT/name).read_bytes()).hexdigest()!=digest:
            raise RuntimeError('Verified evidence changed before report authoring: '+name)
    execution=json.loads((OUT/'run_manifest.json').read_text())
    if execution['status']!='complete': raise RuntimeError('Experiment has not completed')
    core_files={ROOT/'phase2'/name:execution['source_hashes'][name]
                for name in ('analysis.py','decisions.py','experiment.py','plan.py','run.py')}
    core_files.update({ROOT/'src'/name:digest for name,digest in execution['dependency_hashes'].items()})
    for path,digest in core_files.items():
        if hashlib.sha256(path.read_bytes()).hexdigest()!=digest:
            raise RuntimeError('Experimental implementation changed before reporting: '+path.name)
    supplement=json.loads((OUT/'supplement_manifest.json').read_text())
    supplement_check=json.loads((OUT/'supplement_verification.json').read_text())
    sha=lambda path:hashlib.sha256(path.read_bytes()).hexdigest()
    if not supplement_check['passed'] or sha(OUT/'supplement_manifest.json')!=supplement_check['manifest_sha256']:
        raise RuntimeError('Independently verified supplementary accounting is required')
    for name,digest in {**supplement['inputs'],**supplement['outputs']}.items():
        if sha(OUT/name)!=digest: raise RuntimeError('Supplementary evidence changed: '+name)
    if sha(ROOT/'phase2/supplement.py')!=supplement['source_sha256'] or sha(ROOT/'phase2/verify_supplement.py')!=supplement_check['verifier_sha256']:
        raise RuntimeError('Supplementary accounting or verification changed')
    value=read('decision_value'); value_pool=read('decision_value_pooled')
    explanations=read('threshold_explanations'); enriched=read('threshold_audit_enriched')
    transfers=read('refit_transfer'); capacity=read('capacity_tradeoffs'); scenario_value=read('prevalence_decision_value')
    locked=json.loads((ROOT/'phase2/locked_plan.json').read_text())
    manifests=json.loads((OUT/'fold_manifest.json').read_text())
    x=e.query("model=='XGBoost' and scheme=='time' and workflow=='fixed'")
    xc=x.query("mechanism=='C'")
    xa=x.query("mechanism=='A' and cost==10")
    cx=c.query("model=='XGBoost' and fold>=0")
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.82)
    sec.left_margin=sec.right_margin=Inches(.92)
    sec.header_distance=sec.footer_distance=Inches(.35)
    for st in doc.styles:
        for border in list(st.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
        st=doc.styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
        fonts=st.element.get_or_add_rPr().rFonts
        for key in list(fonts.attrib):
            if 'Theme' in key: del fonts.attrib[key]
        for script in ['ascii','hAnsi','cs','eastAsia']: fonts.set(qn('w:'+script),'Times New Roman')
        st.font.size=Pt(11 if name=='Normal' else (16 if name=='Title' else 10 if name=='Caption' else 12))
        st.paragraph_format.space_after=Pt(6)
    normal=doc.styles['Normal']; normal.paragraph_format.line_spacing=1.08
    normal.paragraph_format.widow_control=True
    for name in ['Heading 1','Heading 2']:
        doc.styles[name].font.bold=True
        doc.styles[name].paragraph_format.space_before=Pt(8)
    foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    foot.add_run('Threshold transfer study  |  ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); foot._p.append(field)
    for run in foot.runs: run.font.size=Pt(9)
    md=[]; claims=[]
    def para(text,style=None):
        pr=doc.add_paragraph(text,style)
        if style is None: pr.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        md.append(text+'\n'); return pr
    def heading(text): para(text,'Heading 1')
    def page(text):
        pr=para(text,'Heading 1'); pr.paragraph_format.page_break_before=True
    def table(caption,headers,rows,widths=None):
        para(caption,'Caption'); tab=doc.add_table(rows=1,cols=len(headers)); tab.autofit=True
        for cell,text in zip(tab.rows[0].cells,headers): cell.text=str(text)
        repeat=OxmlElement('w:tblHeader'); tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell,text in zip(tab.add_row().cells,row): cell.text=str(text)
        if widths is not None:
            tab.autofit=False
            for column,width in zip(tab.columns,widths): column.width=Inches(width)
            for row in tab.rows:
                for cell,width in zip(row.cells,widths): cell.width=Inches(width)
        for i,row in enumerate(tab.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for pr in cell.paragraphs:
                    pr.paragraph_format.space_after=Pt(3); pr.paragraph_format.space_before=Pt(3)
                    pr.paragraph_format.line_spacing=1
                    for run in pr.runs: run.font.size=Pt(9); run.bold=(i==0)
                if i==0:
                    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EEEEEE'); cell._tc.get_or_add_tcPr().append(sh)
        # Academic horizontal rules, without a heavy table grid.
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','bottom','insideH']:
            el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'AAAAAA'); borders.append(el)
        tab._tbl.tblPr.append(borders)
        md.extend(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |'])
        md.extend('| '+' | '.join(str(cell).replace('\n','<br>') for cell in row)+' |' for row in rows); md.append('')
    def rng(values):
        lo,hi=min(values),max(values)
        return '%.2f'%lo if lo==hi else '%.2f to %.2f'%(lo,hi)
    def heatmap(filename,panels,caption):
        plt.rcParams.update({'font.family':'DejaVu Sans','font.size':16})
        fig,axes=plt.subplots(1,len(panels),figsize=(11.8,7.7),layout='constrained')
        axes=np.atleast_1d(axes)
        labels=['P%d  r=%d'%(f+1,r) for f in range(4) for r in PLAN['costs']]
        for ax,(title,values,diverging) in zip(axes,panels):
            mat=np.asarray(values).reshape(16,4)
            if diverging:
                limit=max(float(np.abs(mat).max()),1)
                im=ax.imshow(mat,cmap='RdBu_r',vmin=-limit,vmax=limit,aspect='auto')
            else: im=ax.imshow(mat,cmap='Blues',vmin=0,vmax=max(float(mat.max()),1),aspect='auto')
            ax.set_title(title,fontsize=17,pad=12)
            ax.set_xticks(range(4),['1%','5%','10%','20%']); ax.set_xlabel('Reference capacity budget')
            ax.set_yticks(range(16),labels,fontsize=14)
            for row in range(16):
                for col in range(4):
                    v=mat[row,col]; norm=im.norm(v)
                    color='white' if (norm<.16 or norm>.84) and diverging else ('white' if norm>.65 and not diverging else 'black')
                    counts = not diverging or title == 'Refitted minus fixed alerts'
                    text=str(int(v)) if counts else '%.1f'%v
                    ax.text(col,row,text,ha='center',va='center',fontsize=14,color=color)
            for sep in [3.5,7.5,11.5]: ax.axhline(sep,color='black',lw=1.2)
            fig.colorbar(im,ax=ax,shrink=.6,pad=.02)
        fig.savefig(DEST/filename,dpi=220); plt.close(fig)
        pr=doc.add_paragraph(); pr.alignment=WD_ALIGN_PARAGRAPH.CENTER
        pr.add_run().add_picture(str(DEST/filename),width=Inches(5.9))
        para(caption,'Caption'); md.extend(['![%s](%s)'%(caption,filename),''])
    def matrix_from(frame,column):
        return [frame.query('fold==@f and cost==@r and budget==@b')[column].iloc[0]
                for f in range(4) for r in PLAN['costs'] for b in PLAN['budgets']]

    title=para(PLAN['title'],'Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    author=para('Junjie Yang\nMining Engineering, Fuzhou University'); author.alignment=WD_ALIGN_PARAGRAPH.CENTER
    heading('Abstract')
    over=int((xa.excess>0).sum()); c_over=int((xc.excess>0).sum()); silent=int((xc.alerts==0).sum())
    cause=explanations.query("model=='XGBoost' and scheme=='time' and reference_alerts==0")
    cause_counts=cause.no_alarm_reason.value_counts()
    pooled_c=value_pool.query("model=='XGBoost' and workflow=='fixed' and cost==10 and mechanism=='C'")
    abstract=(f'Historical capacity feasibility does not ensure that a seismic warning rule retains its workload or decision value on later records. '
        f'This study separates model fitting, policy selection and testing to evaluate budget-only, cost-only and cost-plus-capacity rules on the UCI Seismic Bumps data. '
        f'Four subsequent record blocks provide 2,063 test records. XGBoost is the prespecified main model; logistic regression and bagged CART provide robustness checks. '
        f'At relative missed-shift cost r=10, pooled cost-plus-capacity loss equaled or exceeded the no-alarm reference by {pooled_c.delta_loss_vs_no_alarm_100.min():.2f} to {pooled_c.delta_loss_vs_no_alarm_100.max():.2f} units per 100 shifts. '
        f'Despite all reference areas containing positive outcomes, its {silent} no-alarm selections comprise {int(cause_counts.get("cost_selects_no_alarm",0))} unconstrained cost optima and {int(cause_counts.get("capacity_changes_to_no_alarm",0))} capacity-induced changes. '
        f'The budget rule exceeded later capacity in {over} of 16 phase–budget settings; cost-plus-capacity exceeded it in {c_over} of 64 phase–budget–cost settings. These dependent settings use different denominators. '
        'Matched-setting comparisons connect reduced capacity excess to changes in detections, missed events and loss. Refitting changes finite-cutoff decisions, while no-alarm rules remain invariant by construction. '
        'Paired block intervals and frozen-rule prevalence scenarios characterize decision value and workload under hypothetical costs. The workflow links historical selection audits to later decision value. Recorded row order is a temporal proxy; the study is retrospective.')
    # Avoid any manuscript number without a matching computed claim source.
    claims.extend([dict(claim='fixed XGBoost A capacity violations',value=over,source='evaluations.csv; time/A/cost10/excess>0'),
                   dict(claim='fixed XGBoost C no alarms',value=silent,source='evaluations.csv; time/C/alerts==0'),
                   dict(claim='fixed XGBoost C capacity violations',value=c_over,source='evaluations.csv; time/C/excess>0'),
                   dict(claim='historical cost optimum no alarms',value=int(cause_counts.get('cost_selects_no_alarm',0)),source='threshold_explanations.csv; XGBoost/time/reference_alerts==0'),
                   dict(claim='historical capacity-induced no alarms',value=int(cause_counts.get('capacity_changes_to_no_alarm',0)),source='threshold_explanations.csv; XGBoost/time/reference_alerts==0')])
    para(abstract)
    para('Keywords: seismic hazard; warning threshold; relative cost; inspection capacity; record-order validation.')
    heading('1 Research question and data')
    para('A warning score becomes an inspection decision through a threshold that trades missed hazardous shifts against false alerts. A cutoff selected under historical costs and capacity may trigger a different workload on later records. This study asks whether such rules reduce loss relative to no alarms, how their inspection demand transfers, and how model refitting changes decisions at the same cutoff. An alert identifies a candidate shift for inspection; the labels do not measure inspection effectiveness or avoided accidents. The contribution is auditable decision analysis that links historical selection reasons to later loss, hazardous-shift coverage and capacity excess, and evaluates refitting at the same numerical cutoff.')
    para('Mining cold-start research addresses limited local history at new sites [10]; microseismic reviews distinguish waveform detection and location from forecasting [11]. Here, cost-sensitive loss [5] and inspection budgets connect shift forecasts to decisions. We jointly evaluate transferred loss, workload and refitting, with historical audits explaining no-alarm choices. This links the selection rationale to later outcomes on the same test records.')
    para('UCI describes eight-hour shift summaries and a next-shift target indicating a seismic bump above 10⁴ J [1]. We retain the first-stage 2,578-row mirror, with 170 positives, and its full 17-column design. The official 2,584-row file contains six duplicate occurrences removed in the mirror; the first-stage source mapping identifies their retained first occurrences. Recorded row order is the temporal proxy because timestamps and operation identifiers are unavailable.')
    para('The analysis separates three questions: how cost and capacity determine historical selection, how the selected rules perform in later blocks, and how their decisions change after model refitting with the same cutoff. The protocol was locked before second-stage fitting. Because the cohort had already been inspected in stage one, the study is exploratory and retrospective. The previously viewed 774-row holdout provides supplementary evidence and overlaps part of the four-block cohort.')

    page('2 Historical selection and transfer design')
    rows=[]; prevalence_changes=[]
    for m in manifests:
        if m['model']!='XGBoost' or m['scheme']!='time': continue
        rr=p[(p.model=='XGBoost')&(p.scheme=='time')&(p.fold==m['fold'])]
        reference_row=enriched[(enriched.model=='XGBoost')&(enriched.scheme=='time')&(enriched.fold==m['fold'])].iloc[0]
        test_prevalence=float(rr.y.mean())
        prevalence_changes.append('phase %d: %.2f%% → %.2f%%'%(m['fold']+1,100*reference_row.reference_prevalence,100*test_prevalence))
        rows.append([m['fold']+1,len(m['fit_rows']),len(m['reference_rows']),int(reference_row.reference_positives),
            '%.2f%%'%(100*reference_row.reference_prevalence),len(m['test_rows']),int(rr.y.sum()),'%.2f%%'%(100*test_prevalence),'%d / %d'%(m['depth'],m['rounds'])])
    table('Table 1. Disjoint areas, reference/test prevalence and training-only XGBoost choices',['Phase','Fit n','Ref n','Ref +','Ref %','Test n','Test +','Test %','Depth / rounds'],rows)
    para('For each outer historical prefix, the first 80% is the model-fitting area and the final 20% is the policy-reference area. XGBoost [2] selects depth {2, 3} and rounds {80, 160} on a further record-order 80/20 split inside the fitting area, using average precision; exact ties select the first grid entry. Learning rate is 0.05, minimum child weight 10 and L2 penalty 5, with full row and column sampling, two threads and seed 7. If that internal reference has no positives, negative Brier score is used. LR uses balanced class weights and penalty 1; bagged CART uses 60 unweighted trees, depth 6, minimum leaf size 20 and seed 7. Configurations are intentionally asymmetric and retained from stage one.')
    para('The fixed model is trained on the fitting area. Reference scores and labels determine a cutoff, which is frozen for the next test block. The refitted comparison uses the already selected parameters to fit all outer historical rows, including the policy-reference area, then applies the identical numerical cutoff to the identical test rows. Thresholds are selected using fixed-model reference scores and applied to raw scores in both workflows. Test labels enter only the subsequent evaluation.')
    para('Stage one [3] transfers a reference cutoff to an outer model refitted on all history; here fixed-model transfer is primary and refitting is a separate comparison. Stage-one XGBoost reference models use a separate parameter search, and budget cutoffs exclude boundary ties. Stage two reuses fitting-area parameters and canonical whole-group cutoffs. Cross-report alert counts require matched models, parameter selection, workflows and threshold definitions.')
    table('Table 2. Rules on the historical reference area',['Rule','Objective','Constraint'],[
        ['A budget','Maximum whole-group alert count','Alerts ≤ floor(B Nref)'],
        ['B cost','Minimum r FN + FP','None'],
        ['C cost plus capacity','Minimum r FN + FP','Alerts ≤ floor(B Nref)']])
    para('Budgets are 1%, 5%, 10% and 20%; relative penalties r for missing a hazardous shift are 5, 10, 20 and 50, with false-alarm cost 1. A is selected once per budget, B once per cost, and C once per combination. Alerts satisfy s ≥ τ. Tied scores move together. The cutoff is the minimum included reference score, except +∞ for no alarms and −∞ for all alarms. Equal-loss B/C decisions select fewer alerts. Zero capacity slots force A/C to +∞; without reference positives B/C select no alarms, while A remains budget-only.')
    para('The no-alarm baseline is always evaluated. Batch Top-k selects the largest whole-tie test-score set within floor(B Ntest), without test labels, and can underfill capacity. This retrospective ranking reference requires scores for the complete test batch; its capacity is enforced on that batch rather than inherited from a historical cutoff.')
    para('Primary loss is L100 = 100(r FN + FP)/N. Counts, recall, precision, alert rate and excess max(0, Alerts − floor(B N)) accompany it. Costs are hypothetical relative weights, not monetary estimates. Pooled slots and excess sum phase-specific quantities without offsetting excess in one phase against spare capacity in another.')

    page('2.1 Decision assumptions and evaluation scope')
    table('Table 2a. Decision assumptions and evaluation scope',['Item','Definition and evaluation scope'],[
        ['Alert object','One dataset shift record flagged as hazardous for the next shift; recorded row order is a temporal proxy.'],
        ['Capacity unit','One flagged record represents one assumed inspection slot. Historical slots are floor(B Nref); later slots are floor(B Ntest), separately for each block. Inspection duration and staffing are not measured.'],
        ['TP / FP / FN','TP: flagged hazardous shift. FP: flagged non-hazardous shift. FN: unflagged hazardous shift. TP counts label coverage, not confirmed inspections or avoided accidents.'],
        ['Relative loss','L = r FN + FP assigns missed hazardous shifts cost r and false alerts cost 1. It includes no separate TP inspection cost, accident-severity cost, delay cost or monetary estimate.'],
        ['Excess alerts','Excess = max(0, Alerts − floor(B Ntest)). All score-based alerts remain in TP/FP/FN and loss. Excess is recorded as demand; alerts are not truncated, queued or assigned an additional loss penalty.'],
        ['Pooled accounting','Sum counts, slots and excess over the four later blocks; divide total loss by total records to obtain loss per 100. Spare slots in one block do not offset excess in another.'],
        ['Operational evidence','Actual inspection execution and accident prevention are not observed. Capacity and relative costs are explicit hypothetical decision settings.']],widths=[1.45,4.98])
    para('These assumptions define how frozen warning rules create inspection demand. The evaluation connects hazardous-shift coverage to relative loss and workload; excess demand remains in the alert counts. Detections count hazardous shifts flagged, while dispatch, queueing and safety interventions are outside the evaluated process.')

    page('3 Decision value relative to no alarms')
    pooled_main=value_pool.query("model=='XGBoost' and workflow=='fixed' and cost==10")
    vrows=[]
    for mechanism in ['none','A','B','C']:
        for _,row in pooled_main[pooled_main.mechanism==mechanism].sort_values('budget').iterrows():
            vrows.append([mechanism,'—' if pd.isna(row.budget) else '%d%%'%round(100*row.budget),
                '%.2f'%row.loss100,'%.2f'%row.delta_loss_vs_no_alarm_100,int(row.alerts),int(row.tp),int(row.fn)])
    table('Table 3. Pooled fixed XGBoost decision value at r=10',['Rule','Budget','L100','Δ vs none','Alerts','TP','FN'],vrows)
    para('The no-alarm rule provides an explicit decision reference: Lnone,100 = 100 r(TP + FN)/N. Relative loss is ΔLnone,100 = 100(FP − rTP)/N. Negative values indicate lower loss under the stated hypothetical cost. A and C budgets are historical selection limits; B has no capacity constraint. The cohort contains 2,063 shifts and 88 hazardous outcomes.')
    stage_rows=[]
    for fold in range(4):
        part=value.query("model=='XGBoost' and scheme=='time' and workflow=='fixed' and cost==10 and fold==@fold")
        stage_rows.append([fold+1,'%.2f'%part[part.mechanism=='none'].loss100.iloc[0],
            rng(part[part.mechanism=='A'].delta_loss_vs_no_alarm_100),
            '%.2f'%part[part.mechanism=='B'].delta_loss_vs_no_alarm_100.iloc[0],
            rng(part[part.mechanism=='C'].delta_loss_vs_no_alarm_100)])
    table('Table 4. Phase-specific loss relative to no alarms at r=10',['Phase','No-alarm L100','A Δ range','B Δ','C Δ range'],stage_rows)
    para('At this cost, pooled C loss equals the no-alarm loss at the 1% budget and is higher at the remaining budgets. Historical optimization is compatible with this outcome: reference labels determine the optimum, whereas later false alarms and detected hazards determine transferred value. At the 5%, 10% and 20% budgets, pooled C produces 24, 27 and 32 false alarms for 1, 1 and 2 detections, respectively; each exceeds the corresponding rTP benefit at r=10. Phase-specific results show where these contributions arise rather than attributing the pooled result to every stage.')
    para('A lower loss than another warning rule does not establish value over no alarms. C can reduce workload and loss relative to A while retaining positive ΔLnone,100. No alarms incur the cost of every missed hazardous shift and detect none; their role here is a loss reference, not an operational recommendation. Complete files retain all costs, models, workflows and the descriptive holdout; r=10 is the common display scale.')

    page('4 Budget thresholds and future capacity')
    arows=[]
    for _,row in xa.sort_values(['fold','budget']).iterrows():
        audit=a[(a.model=='XGBoost')&(a.scheme=='time')&(a.fold==row.fold)&(a.rule==row.rule)].iloc[0]
        arows.append([int(row.fold)+1,'%d%%'%round(100*row.budget),int(audit.reference_alerts),int(row.alerts),int(row.tp),int(row.fn),int(row.excess),'%.2f'%row.loss100])
    table('Table 5. Fixed XGBoost budget rule A across all phases and budgets',['Phase','Budget','Ref alerts','Test alerts','TP','FN','Excess','L100 r=10'],arows)
    para('The loss column uses r=10 as a common display scale; A itself never uses the cost ratio or labels to choose a cutoff. Each phase contains 515 or 516 test shifts. All 16 reference choices satisfied their own historical capacity limits.')
    para(f'Future capacity was exceeded in {over} of the 16 combinations. In phase 2, the 20% historical budget transferred to 254 alerts among 515 shifts, with 151 alerts beyond the 103 available slots. In phase 1, that same budget produced 21 alerts among 516 shifts, below the 103-slot limit. A budget therefore constrains historical selection rather than subsequent workload. Tied-score exclusions and changes in the score distribution can both separate the intended budget from its realized rate.')
    para('The workload difference accompanies a detection tradeoff. Phase 2 at the 20% budget detected 6 of 10 hazardous shifts and produced 248 false alarms; phase 1 detected 2 of 36 hazardous shifts with 19 false alarms. Reading these counts together shows how the same historical budget can yield both underused inspection capacity and substantial capacity excess in later stages.')
    para('The electronic evaluation table reports recall, precision and alert rate for every cost and workflow. Batch Top-k supplies a complementary reference: it enforces capacity using the observed test-batch scores, while the frozen threshold carries forward a historical selection. Their workload difference therefore reflects the information available when each rule is applied.')

    page('5 Cost and capacity change the selected decision')
    ca=cx.query("comparison=='C-A' and workflow=='fixed'")
    cb=cx.query("comparison=='C-B' and workflow=='fixed'")
    heatmap('loss_contrasts.png',[
        ('C minus A loss per 100',matrix_from(ca,'loss_delta100'),True),
        ('C minus B loss per 100',matrix_from(cb,'loss_delta100'),True)],
        'Figure 1. Fixed XGBoost loss contrasts for every phase, cost and capacity budget. Negative values favor C on the stated relative loss. Panels use separate color scales; cell values are rounded to one decimal.')
    para(f'Across the 64 phase–budget–cost combinations, C−A ranged from {rng(ca.loss_delta100)} relative loss units per 100 shifts and C−B from {rng(cb.loss_delta100)}. Strategy preference depends on the record stage and assumed missed-event cost: the same capacity-constrained rule can reduce loss in one setting and increase it in another.')
    para('In phase 2, both B and C selected no alarms for every cost, so C−B is zero. The same decision can lower loss relative to A when r is modest and increase it when missed events carry a higher assumed cost. In phase 1, C reduces inspection work compared with the large alert set chosen by B at higher costs, while accepting more missed hazardous shifts. The positive C−B difference is the observed loss tradeoff after transferring the capacity-constrained rule, not a guaranteed penalty of capacity constraints.')
    paired=capacity.query("model=='XGBoost' and cohort=='pooled' and workflow=='fixed' and cost==10 and budget==0.2").iloc[0]
    para(f'At r=10 and the 20% reference budget, pooled C−A reduced excess by {int(-paired.excess_delta)} alerts but also reduced detections from {int(paired.A_tp)} to {int(paired.C_tp)}, increasing misses from {int(paired.A_fn)} to {int(paired.C_fn)}. Relative loss fell by {-paired.loss100_delta:.2f} per 100 shifts. This is a workload–detection tradeoff; Table A1 presents each phase and budget, and Section 3 gives the separate no-alarm comparison.')
    para('On the reference area, C has loss at least as large as B because its feasible candidate set is a subset; every saved selection satisfies this ordering. After transfer, the ordering can reverse. Historical optimization and subsequent evaluation thus answer distinct questions: the first identifies the best feasible reference decision, and the second measures how that decision performs on later records.')

    page('6 Workload and historical no-alarm selection')
    heatmap('workload.png',[
        ('C test alerts',matrix_from(xc,'alerts'),False),
        ('C alerts beyond capacity',matrix_from(xc,'excess'),False)],
        'Figure 2. Workload and capacity excess for the same fixed XGBoost C policies shown in Figure 1. Zero alerts and zero excess are distinct outcomes. Counts refer to whole test blocks, not rates.')
    para(f'C produced no alarms in {silent} of 64 settings, all following a historical no-alarm selection. Its {c_over} excess-capacity settings and A’s {over} use different denominators; these are grid summaries, not directly comparable independent violation rates. The matched C−A table in the appendix pairs capacity excess with detection and loss.')
    reason_rows=[]
    for fold in range(4):
        part=cause[cause.fold==fold]
        reason_rows.append([fold+1,int((part.no_alarm_reason=='no_reference_positives').sum()),
            int((part.no_alarm_reason=='cost_selects_no_alarm').sum()),int((part.no_alarm_reason=='capacity_changes_to_no_alarm').sum()),
            int(part.reference_capacity_zero.sum()),int(part.no_alarm_tied_with_other_optimum.sum())])
    table('Table 6. Historical reasons for all 40 fixed XGBoost C no-alarm settings',['Phase','No ref +','Cost optimum','Capacity change','Zero slots','Other optimal tie'],reason_rows)
    para('All reference areas contain hazardous shifts: 28 no-alarm settings are cost optima and 12 are capacity-induced changes; none has zero slots or an optimal tie. In phase 3 at r=50, B selects alarms while C selects none at every budget. Reasons use reference records only.')
    audit_rows=[]
    ac=a.query("model=='XGBoost' and scheme=='time' and mechanism=='C'")
    for f,q in ac.groupby('fold'):
        audit_rows.append([int(f)+1,int(q.capacity_binding.sum()),rng(100*q.loss_gap/q.reference_n)])
    table('Table 7. Reference capacity and standardized selection margin',['Phase','Unconstrained cost optimum exceeds reference capacity','Next-decision ΔL100 range'],audit_rows)
    para('The capacity column counts infeasible B optima; C utilization is stored separately. The margin is 100(Lnext − Lbest)/Nref for a second distinct feasible alarm set, undefined if absent and zero for ties. A small margin indicates weak reference advantage, not threshold instability. The enriched audit retains alerts and slots.')

    page('7 Transferring the same cutoff to a refitted model')
    rc=cx[(cx.comparison=='Refitted-Fixed')&cx.left_rule.str.startswith('C')]
    heatmap('refit_contrasts.png',[
        ('Refitted minus fixed loss per 100',matrix_from(rc,'loss_delta100'),True),
        ('Refitted minus fixed alerts',matrix_from(rc,'alert_delta'),True)],
        'Figure 3. C policy changes after fitting on all outer history, with parameters and numerical thresholds held fixed. Negative loss differences favor refitting; negative alert differences mean lower inspection demand. Alert-count changes in the second panel are labelled as integers.')
    para(f'For C, refitted-minus-fixed loss ranged from {rng(rc.loss_delta100)} per 100 shifts and alert changes from {int(rc.alert_delta.min())} to {int(rc.alert_delta.max())} per test block. Phase 2 and phase 3 C decisions remained no-alarm under both workflows. Phase 1 and phase 4 show that the same numerical cutoff can produce different workloads and loss after additional model fitting.')
    main_transfer=transfers.query("model=='XGBoost' and scheme=='time'")
    refit_rows=[]
    for mechanism,part in main_transfer.groupby('mechanism',sort=True):
        distinct=part[['fold','rule','threshold_kind']].drop_duplicates()
        finite=part[part.threshold_kind=='finite']
        refit_rows.append([mechanism,int((distinct.threshold_kind=='finite').sum()),int((distinct.threshold_kind=='no_alarm').sum()),
            int((distinct.threshold_kind=='all_alarm').sum()),rng(finite.loss100_delta),'%d to %d'%(finite.alerts_delta.min(),finite.alerts_delta.max())])
    table('Table 8. XGBoost refitting across all historically finite and extreme rules',['Rule','Finite','+∞','−∞','Finite ΔL100 range','Finite Δ alerts'],refit_rows)
    para('Counts are distinct phase–rule settings: A = phase × budget (4 × 4 = 16); B = phase × cost (4 × 4 = 16); C = phase × budget × cost (4 × 4 × 4 = 64). A is reused across costs. Ranges evaluate finite settings at all four costs. Extreme rules stay unchanged by construction, providing no evidence of finite-cutoff robustness. The appendix retains every A/B/C contrast.')
    para('Refitting changes the model to which the historical rule is applied. Adding historical records, including the former policy-reference area, can alter both the fitted decision function and its score scale. The comparison measures their combined workflow difference rather than isolating a causal contribution. Since selection and application use raw scores throughout, the results specifically describe numerical-cutoff transfer across fitted models.')
    para('Within the observed settings, C exhibited a narrower finite-threshold alert-change range. This describes the tested rule grids. Batch Top-k remains outside the same-cutoff comparison because each complete test-score batch determines its own threshold.')

    page('8 Paired uncertainty and model robustness')
    para('For each prespecified loss contrast, 2,000 paired moving-block samples use block length 32 separately within phases [4], without cross-phase or circular blocks. Both policies use the same sampled rows, and pooled resampling preserves phase sizes. Percentile 95% intervals condition on the fitted models, chosen thresholds and observed cohort, excluding tuning, fitting and policy-selection uncertainty. Identical paired decisions yield degenerate zero intervals.')
    ci_rows=[]
    for name,frame in [('C−A',ca),('C−B',cb),('Refitted−Fixed C',rc)]:
        for f,q in frame.groupby('fold'):
            ci_rows.append([name,int(f)+1,rng(q.loss_delta100),int((q.ci_high<0).sum()),int((q.ci_low>0).sum()),int(((q.ci_low<=0)&(q.ci_high>=0)).sum())])
    table('Table 9. Phase-specific interval directions across all 16 settings per contrast',['Contrast','Phase','Point range','CI below 0','CI above 0','CI contains 0'],ci_rows)
    para('The full file reports the point and both interval endpoints for each individual setting, alongside alert and excess differences. Table 9 summarizes the grid rather than averaging unlike costs; interval counts are descriptive and have no multiple-comparison adjustment. B−A is retained as a secondary contrast.')
    interval_examples=ca[(ca.cost==10)&(ca.budget==.2)&ca.fold.isin([0,1])].sort_values('fold')
    if len(interval_examples)!=2: raise RuntimeError('Missing matched illustrative intervals')
    first,second=[row for _,row in interval_examples.iterrows()]
    para(f'For illustration, hold XGBoost, r=10 and budget 20% fixed. Phase 1 C−A is {first.loss_delta100:+.2f} per 100 shifts '
         f'(95% CI [{first.ci_low:.2f}, {first.ci_high:.2f}]); phase 2 is {second.loss_delta100:+.2f} '
         f'([{second.ci_low:.2f}, {second.ci_high:.2f}]). The positive phase-1 point estimate has an interval containing zero; '
         'the phase-2 interval is entirely negative. These existing intervals illustrate stage-specific uncertainty without selecting an operating rule.')
    for _,row in interval_examples.iterrows():
        claims.append(dict(claim='illustrative fixed XGBoost C-A interval',fold=int(row.fold),cost=10,budget=.2,
            point=float(row.loss_delta100),ci_low=float(row.ci_low),ci_high=float(row.ci_high),source='comparisons.csv'))
    robust=[]
    for model in ['LR','CART']:
        for f in range(4):
            q=e.query("model==@model and scheme=='time' and workflow=='fixed' and mechanism=='C' and fold==@f")
            diff=c.query("model==@model and fold==@f and comparison=='C-A' and workflow=='fixed'")
            robust.append([model,f+1,rng(diff.loss_delta100),'%d–%d'%(q.alerts.min(),q.alerts.max()),int((q.excess>0).sum()),int((q.alerts==0).sum())])
    table('Table 10. Robustness models across the same 16 C settings in each phase',['Model','Phase','C−A L100 range','Alert range','Excess policies','No-alarm policies'],robust)
    para('LR and CART extend the policy comparison across retained baseline models using identical reference boundaries, selection rules and test rows. Their stage-dependent loss and workload ranges show how the tradeoffs vary across fitted score distributions. Model configurations follow the asymmetric first-stage design described in Section 2. Full fixed/refitted and no-alarm comparisons appear in the electronic appendix.')

    page('9 Frozen prevalence scenarios and decision implications')
    para('The scenario changes hazardous-shift prevalence π to 2%, 5%, 10% or 15%, holding policies and empirical class-conditional error rates fixed. Expected loss is L100(π) = 100[r π FNR + (1 − π) FPR]; expected alert rate is π(1 − FNR) + (1 − π)FPR. Cutoffs remain frozen. These prevalence-only expectations exclude within-class changes and do not estimate capacity-exceedance probability.')
    s=scenario_value.query("model=='XGBoost' and scheme=='time' and fold==-1 and workflow=='fixed' and rule.str.startswith('C')")
    srows=[]
    for r in PLAN['costs']:
        row=[r]
        for pi in PLAN['scenarios']:
            part=s[(s.cost==r)&(s.scenario_prevalence==pi)]
            row.append(rng(part.delta_loss_vs_no_alarm_100)+'\n'+rng(100*part.expected_alert_rate)+'%')
        srows.append(row)
    table('Table 11. Pooled prior-shift sensitivity: relative C loss and expected workload',['Cost r','π 2%','π 5%','π 10%','π 15%'],srows)
    para('Each cell ranges across four budgets: first ΔLnone,100(π) = Lpolicy,100(π) − 100rπ, then expected alert percentage. Negative differences mean lower hypothetical loss. These marginal endpoints need not share a budget; Table A5 pairs values by budget at r=10. Absolute losses and per-phase expectations remain in the appendices.')
    hold=e.query("scheme=='holdout' and mechanism=='C'")
    hrows=[]
    for (model,w),q in hold.groupby(['model','workflow'],sort=False):
        hrows.append([model,w,rng(q.loss_delta_no_alarm100),'%d–%d'%(q.alerts.min(),q.alerts.max()),int((q.excess>0).sum())])
    table('Table 12. Previously viewed holdout and all 16 C settings',['Model','Workflow','L100 minus no alarm range','Alert range','Excess policies'],hrows)
    heading('10 Discussion and conclusions')
    sparse=enriched.query("model=='XGBoost' and scheme=='time' and fold==1").iloc[0]
    para('Reference-to-test hazardous-shift prevalence changes were '+ '; '.join(prevalence_changes)+'. '
         f'Phase 2 reference choices rest on {int(sparse.reference_positives)} hazardous outcome among {int(sparse.reference_n)} records. '
         'This sparse reference and the prevalence differences provide plausible context for transfer behavior. Class-conditional score distributions may also change; '
         'the design does not identify prevalence shift as its sole cause. Section 9 uses the label-shift assumption [6] to vary prevalence while freezing class-conditional rates; it does not estimate or correct the actual shift.')
    para('Later decision value is assessed jointly through loss, hazardous-shift coverage and workload. At r=10, C can reduce loss relative to A without improving pooled loss relative to no alarms. Its 40 historical no-alarm choices comprise 28 cost optima and 12 capacity-induced changes despite positive reference outcomes. Matched C−A results quantify the accompanying detections, misses and loss. Refitting changes finite-cutoff decisions; invariant no-alarm rules do not establish finite-threshold stability. Selection audits and paired transfer comparisons connect the historical reason for each cutoff to its later loss and inspection demand.')
    para('Historical selection, later capacity and model updating require separate evidence. Cross-longwall transfer [7] and new-site cold start [10] concern identified locations; our cutoff comparisons concern later record blocks. Structured validation distinguishes interpolation from extrapolation [8]. Costs and capacity remain hypothetical. Further evidence should cover complete fitting/selection uncertainty and a preregistered rule on an independent, timestamped working face with recorded inspection outcomes.')
    page('Appendix A Detection, capacity and pooled evidence')
    capacity_rows=[]
    paired_phase=capacity.query("model=='XGBoost' and scheme=='time' and cohort=='phase_or_holdout' and workflow=='fixed' and cost==10")
    for _,row in paired_phase.sort_values(['fold','budget']).iterrows():
        capacity_rows.append([int(row.fold)+1,'%d%%'%round(100*row.budget),int(row.alerts_delta),int(row.excess_delta),
            int(row.tp_delta),int(row.fn_delta),'%.2f'%row.loss100_delta])
    table('Table A1. Matched C−A capacity, detection and loss differences at r=10',['Phase','Budget','Δ alerts','Δ excess','Δ TP','Δ FN','Δ L100'],capacity_rows)
    para('Every difference is C minus A on the same test rows, with the same historical budget and assumed cost. Negative excess means fewer over-capacity alerts; negative TP means fewer detections, and ΔFN = −ΔTP. Loss and detection differences retain their signs rather than treating fewer alarms as intrinsically better. The complete capacity_tradeoffs file includes both workflows, all costs, all models, pooled results and the descriptive holdout.')
    pooled=read('pooled_evaluations')
    pooled_x=pooled.query("model=='XGBoost' and workflow=='fixed' and cost==10")
    pooled_rows=[]
    for mechanism in ['none','A','C','TopK']:
        for _,row in pooled_x[pooled_x.mechanism==mechanism].sort_values('budget').iterrows():
            pooled_rows.append([mechanism,'—' if pd.isna(row.budget) else '%d%%'%round(100*row.budget),
                               '%.2f'%row.loss100,int(row.alerts),int(row.tp),int(row.fp),int(row.fn),
                               '—' if pd.isna(row.excess) else str(int(row.excess))])
    table('Table A2. Supplementary pooled fixed XGBoost and batch-reference results at r=10',['Rule','Budget','L100','Alerts','TP','FP','FN','Excess'],pooled_rows)
    para('The pooled cohort has 2,063 records and 88 hazardous shifts. Table A2 supplies a common cost scale for retrospective batch ranking. Capacity excess sums stage-specific counts. A and Top-k selections do not depend on r; Table 3 additionally presents the unconstrained cost rule B and no-alarm loss differences.')
    page('Appendix B Refitting and pooled prior shift')
    phase_refit=[]
    for (fold,mechanism),part in main_transfer.groupby(['fold','mechanism']):
        distinct=part[['rule','threshold_kind']].drop_duplicates()
        finite=part[part.threshold_kind=='finite']
        phase_refit.append([int(fold)+1,mechanism,int((distinct.threshold_kind=='finite').sum()),int((distinct.threshold_kind=='no_alarm').sum()),
            int((distinct.threshold_kind=='all_alarm').sum()),rng(finite.loss100_delta) if len(finite) else '—',
            '%d to %d'%(finite.alerts_delta.min(),finite.alerts_delta.max()) if len(finite) else '—'])
    table('Table A3. Phase-specific XGBoost A/B/C refitting with historical threshold groups',['Phase','Rule','Finite','+∞','−∞','Finite ΔL100 range','Finite Δ alerts'],phase_refit)
    para('Groups follow historical threshold choices and retain zero changes. Extreme rules have invariant decisions. refit_transfer.csv contains all 540 A/B/C contrasts across three models, four phases and the descriptive holdout; comparisons.csv retains the original primary-cohort paired intervals.')
    absolute_rows=[]
    for r in PLAN['costs']:
        absolute_rows.append([r]+[rng(s[(s.cost==r)&(s.scenario_prevalence==pi)].expected_loss100) for pi in PLAN['scenarios']])
    table('Table A4. Pooled expected absolute C loss per 100 across four historical budgets',['Cost r','π 2%','π 5%','π 10%','π 15%'],absolute_rows)
    page('Appendix C Budget-specific prior shift and reproducibility')
    paired_scenarios=[]
    for budget in PLAN['budgets']:
        row=['%d%%'%round(100*budget)]
        for pi in PLAN['scenarios']:
            part=s[(s.cost==10)&(s.budget==budget)&(s.scenario_prevalence==pi)]
            if len(part)!=1: raise RuntimeError('Budget-specific prior-shift record is not unique')
            result=part.iloc[0]
            row.append('%.2f\n%.2f%%'%(result.delta_loss_vs_no_alarm_100,100*result.expected_alert_rate))
        paired_scenarios.append(row)
    table('Table A5. Budget-specific pooled fixed XGBoost C prior-shift results at r=10',['Historical budget','π 2%','π 5%','π 10%','π 15%'],paired_scenarios)
    para('Each cell pairs relative loss per 100 shifts with expected alert percentage for one budget and scenario. Negative differences favor C over no alarms at r=10. All sixteen entries freeze historical rules and class-conditional rates; full files retain every model, cost, workflow and phase.')
    heading('Reproducibility and electronic appendix')
    para(f'The protocol was locked on 2 October 2026 at 19:20:26 UTC; its hash starts {locked["plan_sha256"][:16]}. Independent reproduction replayed 15 model pairs and paired resampling. The results/phase2 directory retains row evidence, threshold candidates, audits, evaluations, contrasts and bootstrap arrays. Infinite thresholds are sentinels; empty cells denote undefined quantities.')
    para('Manifests and checks establish provenance for generated tables and figures. Code, commands and complete evidence: https://github.com/junjie-yang11/seismic-bump-forecasting. Reference metadata were cross-checked with the citation-management tools in Scientific Agent Skills [9].')
    para('The author directed scope, protocol refinement and review priorities. ChatGPT and Codex substantially assisted implementation, execution, verification and writing. CONTRIBUTIONS.md in the repository separates author review from computational checks and links project decisions to evidence.')
    page('References')
    para('[1] Sikora M, Wrobel L. seismic-bumps. UCI Machine Learning Repository; 2010. doi:10.24432/C5W902. https://archive.ics.uci.edu/dataset/266/seismic+bumps.')
    para('[2] Chen T, Guestrin C. XGBoost: A Scalable Tree Boosting System. KDD; 2016. doi:10.1145/2939672.2939785.')
    para('[3] Yang J. Evaluating reliability and explainability in seismic hazard forecasting for underground mine monitoring. Technical report; 2026. Companion stage-one paper: reports/phase1/technical_report.pdf in the project repository.')
    para('[4] Künsch HR. The Jackknife and the Bootstrap for General Stationary Observations. The Annals of Statistics; 1989;17(3):1217–1241. doi:10.1214/aos/1176347265.')
    para('[5] Elkan C. The Foundations of Cost-Sensitive Learning. IJCAI; 2001:973–978. Author manuscript: https://cseweb.ucsd.edu/~elkan/rescale.pdf.')
    para('[6] Lipton ZC, Wang YX, Smola AJ. Detecting and Correcting for Label Shift with Black Box Predictors. ICML; PMLR 80; 2018:3122–3130. https://proceedings.mlr.press/v80/lipton18a.html.')
    para('[7] Sikora M, Wrobel L. Application of rule induction algorithms for analysis of data collected by seismic hazard monitoring systems in coal mines. Archives of Mining Sciences; 2010;55(1):91–114. Author-linked full text: https://www.researchgate.net/publication/281395657.')
    para('[8] Roberts DR et al. Cross-validation strategies for data with temporal, spatial, hierarchical, or phylogenetic structure. Ecography; 2017;40:913–929. doi:10.1111/ecog.02881.')
    para('[9] Kassis T, Agarwal V, He Y, Patel D, Brueckner AM. Scientific Agent Skills: A Library of Procedural Knowledge for Research Agents. arXiv; 2026. doi:10.48550/arXiv.2609.00065. https://arxiv.org/abs/2609.00065.')
    para('[10] Janusz A, Grzegorowski M, Michalak M, Wrobel L, Sikora M, Slezak D. Predicting seismic events in coal mines based on underground sensor measurements. Engineering Applications of Artificial Intelligence; 2017;64:83–94. doi:10.1016/j.engappai.2017.06.002.')
    para('[11] Anikiev D, Birnie C, Waheed Ub, Alkhalifah T, Gu C, Verschuur DJ, Eisner L. Machine learning in microseismic monitoring. Earth-Science Reviews; 2023;239:104371. doi:10.1016/j.earscirev.2023.104371.')
    doc.core_properties.author=''; doc.core_properties.last_modified_by=''; doc.core_properties.title=PLAN['title']
    doc.core_properties.subject='Second-stage retrospective warning decision study'
    doc.save(DEST/'phase2_threshold_transfer_report.docx')
    # Remove inherited Office properties using the existing pure package helper.
    import sys
    sys.path.insert(0,str(ROOT))
    from phase1.generate_report import scrub_report_package
    scrub_report_package(DEST/'phase2_threshold_transfer_report.docx')
    (DEST/'phase2_threshold_transfer_report.md').write_text('\n'.join(md),encoding='utf-8')
    inputs={str(f.relative_to(OUT)):hashlib.sha256(f.read_bytes()).hexdigest() for f in OUT.glob('*.csv')}
    (DEST/'report_manifest.json').write_text(json.dumps(dict(inputs=inputs,verification=verification,supplement_verification=supplement_check,
        computed_claims=claims,plan_sha256=locked['plan_sha256'],generator_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest()),indent=2),encoding='utf-8')
    print('Independent phase-two Word and manuscript generated from verified evidence.')

if __name__=='__main__': build()
