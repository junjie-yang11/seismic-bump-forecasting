"""Independently check Stage 2 reports; failed runs invalidate older certificates."""
import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile


def _interval_overview_rows(document):
    """Locate the interval overview across editions with different table counts."""
    headers = ['Contrast', 'Phase', 'Point range', 'CI below 0',
               'CI above 0', 'CI contains 0']
    matches = [table for table in document.tables
               if table.rows and [cell.text for cell in table.rows[0].cells] == headers]
    if len(matches) != 1:
        raise AssertionError('Expected exactly one interval-overview table')
    return [[cell.text for cell in row.cells] for row in matches[0].rows]


def _run_checks(root, baseline):
    root=root.resolve();out=root/'results/phase2';dest=root/'reports/phase2'
    # Parse current input before optional document imports, including in failure tests.
    man=json.loads((dest/'report_manifest.json').read_text(encoding='utf-8'))
    from docx import Document
    from pypdf import PdfReader
    from lxml import etree
    import pandas as pd
    doc=Document(dest/'phase2_threshold_transfer_report.docx');pdf=PdfReader(dest/'phase2_threshold_transfer_report.pdf')
    md=(dest/'phase2_threshold_transfer_report.md').read_text(encoding='utf-8')
    checks=0
    def check(test,message):
        nonlocal checks
        if not test: raise AssertionError(message)
        checks+=1
    sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
    for file,digest in man['inputs'].items():check(sha(out/file)==digest,'CSV source '+file)
    check(sha(root/'phase2/report.py')==man['generator_sha256'],'current report generator')
    run=json.loads((out/'run_manifest.json').read_text())
    for file,digest in run['first_stage_hashes'].items():check(sha(root/file)==digest,'first stage unchanged '+file)
    for file in ['analysis.py','decisions.py','experiment.py','plan.py','run.py']:
        check(sha(root/'phase2'/file)==run['source_hashes'][file],'model/selection implementation unchanged')
    check(13 <= len(pdf.pages) <= 18,'expected report page range');check(len(doc.tables)==20,'twenty complete tables');check(len(doc.inline_shapes)==3,'three figures')
    check(not doc.core_properties.author and not doc.core_properties.last_modified_by,'clean Word metadata')
    check('/Author' not in (pdf.metadata or {}),'clean PDF metadata')
    for para in doc.paragraphs:
        if para.text.strip():check(para.text in md,'Word/manuscript paragraph agreement')
    for table in doc.tables:
        for row in table.rows:
            check('| '+' | '.join(cell.text.replace('\n','<br>') for cell in row.cells)+' |' in md,'Word/manuscript table agreement')
    texts=[page.extract_text() for page in pdf.pages]
    pdftext='\n'.join(texts)
    numbers=lambda t:set(re.findall(r'-?\d+\.\d+',t.replace('−','-')))
    for number in numbers('\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)):
        check(number in numbers(pdftext),'PDF table number '+number)
    for n,text in enumerate(texts):check(len(text)>150,'no blank/orphan page '+str(n+1))
    check('References' in texts[-1] and '[1]' in texts[-1] and '[2]' in texts[-1] and '[3]' in texts[-1] and '[4]' in texts[-1],'references together')
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with ZipFile(dest/'phase2_threshold_transfer_report.docx') as package:
        for name in ['word/styles.xml','word/document.xml']:
            tree=etree.fromstring(package.read(name));check(not tree.xpath('//w:pBdr',namespaces=ns),'no title borders')
        check(not any('comments' in n or n.startswith('customXml/') for n in package.namelist()),'no private comments/custom stores')
    read=lambda name:pd.read_csv(out/(name+'.csv'))
    assumption_tables=[t for t in doc.tables if [c.text for c in t.rows[0].cells]==['Item','Definition and evaluation scope']]
    check(len(assumption_tables)==1,'one decision-assumption table')
    check('Actual inspection execution and accident prevention are not observed.' in md,'operational scope explicit')
    new_headers=[['Evidence','Stage 1 assessment','Stage 2 assessment'],
                 ['Historical budget','TP','FP','Alerts','Equality cost r*']]
    evidence_tables=[t for t in doc.tables if [c.text for c in t.rows[0].cells] not in
                     [['Item','Definition and evaluation scope']]+new_headers]
    getrows=lambda index:[[cell.text for cell in row.cells] for row in evidence_tables[index].rows][1:]
    def rng(values):
        lo,hi=min(values),max(values);return '%.2f'%lo if lo==hi else '%.2f to %.2f'%(lo,hi)
    value=read('decision_value');pooled=read('decision_value_pooled');audit=read('threshold_audit_enriched')
    pred=read('predictions'); folds=json.loads((out/'fold_manifest.json').read_text()); rows=[]
    for m in folds:
        if m['model']!='XGBoost' or m['scheme']!='time': continue
        reference=audit[(audit.model=='XGBoost')&(audit.scheme=='time')&(audit.fold==m['fold'])]
        ref=reference.iloc[0]; test=pred[(pred.model=='XGBoost')&(pred.scheme=='time')&(pred.fold==m['fold'])]
        rows.append([str(m['fold']+1),str(len(m['fit_rows'])),str(len(m['reference_rows'])),str(int(ref.reference_positives)),
            '%.2f%%'%(100*ref.reference_positives/ref.reference_n),str(len(m['test_rows'])),str(int(test.y.sum())),
            '%.2f%%'%(100*test.y.mean()),'%d / %d'%(m['depth'],m['rounds'])])
    check(getrows(0)==rows,'Table 1 reference/test counts and prevalence')
    old=Document(baseline/'reports/phase2/phase2_threshold_transfer_report.docx')
    check(_interval_overview_rows(doc)==_interval_overview_rows(old),'Table 9 preserved exactly')
    comparisons=read('comparisons')
    examples=comparisons.query("model=='XGBoost' and comparison=='C-A' and workflow=='fixed' and cost==10 and budget==0.2 and fold in [0,1]").sort_values('fold')
    check(len(examples)==2,'two existing matched illustrative intervals')
    for _,r in examples.iterrows():
        check('%.2f'%r.loss_delta100 in md,'illustrative interval point')
        check('[%.2f, %.2f]'%(r.ci_low,r.ci_high) in md,'illustrative interval endpoints')
    check('despite all reference areas containing positive outcomes' in md.lower(),'abstract explains positive reference outcomes')
    check('does not identify prevalence shift as its sole cause' in md,'prevalence discussion bounded')
    check('not a guaranteed penalty' in md,'capacity transfer interpretation bounded')
    rows=[]
    q=pooled.query("model=='XGBoost' and workflow=='fixed' and cost==10")
    for mechanism in ['none','A','B','C']:
        for _,row in q[q.mechanism==mechanism].sort_values('budget').iterrows():
            rows.append([mechanism,'—' if pd.isna(row.budget) else '%d%%'%round(100*row.budget),'%.2f'%row.loss100,'%.2f'%row.delta_loss_vs_no_alarm_100,str(int(row.alerts)),str(int(row.tp)),str(int(row.fn))])
    check(getrows(2)==rows,'Table 3 direct reconstruction')
    rows=[]
    for f in range(4):
        q=value.query("model=='XGBoost' and scheme=='time' and workflow=='fixed' and cost==10 and fold==@f")
        rows.append([str(f+1),'%.2f'%q[q.mechanism=='none'].loss100.iloc[0],rng(q[q.mechanism=='A'].delta_loss_vs_no_alarm_100),'%.2f'%q[q.mechanism=='B'].delta_loss_vs_no_alarm_100.iloc[0],rng(q[q.mechanism=='C'].delta_loss_vs_no_alarm_100)])
    check(getrows(3)==rows,'Table 4 direct reconstruction')
    reason=read('threshold_explanations').query("model=='XGBoost' and scheme=='time' and reference_alerts==0")
    rows=[]
    for f in range(4):
        q=reason[reason.fold==f]
        rows.append([str(f+1),str(int((q.no_alarm_reason=='no_reference_positives').sum())),str(int((q.no_alarm_reason=='cost_selects_no_alarm').sum())),str(int((q.no_alarm_reason=='capacity_changes_to_no_alarm').sum())),str(int(q.reference_capacity_zero.sum())),str(int(q.no_alarm_tied_with_other_optimum.sum()))])
    check(getrows(5)==rows,'Table 6 exhaustive historical reasons')
    q=audit.query("model=='XGBoost' and scheme=='time' and mechanism=='C'")
    rows=[[str(int(f)+1),str(int(part.capacity_binding.sum())),rng(part.loss_gap_100)] for f,part in q.groupby('fold')]
    check(getrows(6)==rows,'Table 7 capacity/margin reconstruction')
    transfer=read('refit_transfer').query("model=='XGBoost' and scheme=='time'");rows=[]
    for mechanism,q in transfer.groupby('mechanism'):
        distinct=q[['fold','rule','threshold_kind']].drop_duplicates();finite=q[q.threshold_kind=='finite']
        rows.append([mechanism,str(int((distinct.threshold_kind=='finite').sum())),str(int((distinct.threshold_kind=='no_alarm').sum())),str(int((distinct.threshold_kind=='all_alarm').sum())),rng(finite.loss100_delta),'%d to %d'%(finite.alerts_delta.min(),finite.alerts_delta.max())])
    check(getrows(7)==rows,'Table 8 finite/extreme reconstruction')
    trade=read('capacity_tradeoffs').query("model=='XGBoost' and scheme=='time' and cohort=='phase_or_holdout' and workflow=='fixed' and cost==10")
    rows=[[str(int(r.fold)+1),'%d%%'%round(100*r.budget),str(int(r.alerts_delta)),str(int(r.excess_delta)),str(int(r.tp_delta)),str(int(r.fn_delta)),'%.2f'%r.loss100_delta] for _,r in trade.sort_values(['fold','budget']).iterrows()]
    check(getrows(12)==rows,'Table A1 matched detection and workload reconstruction')
    scen=read('prevalence_decision_value').query("model=='XGBoost' and scheme=='time' and fold==-1 and workflow=='fixed' and rule.str.startswith('C')")
    rows=[]
    for r in [5,10,20,50]:
        row=[str(r)]
        for pi in [.02,.05,.10,.15]:
            q=scen[(scen.cost==r)&(scen.scenario_prevalence==pi)];row.append(rng(q.delta_loss_vs_no_alarm_100)+'\n'+rng(100*q.expected_alert_rate)+'%')
        rows.append(row)
    check(getrows(10)==rows,'Table 11 loss and workload reconstruction')
    original_scen=read('prevalence_scenarios').query("model=='XGBoost' and scheme=='time' and fold==-1 and workflow=='fixed' and cost==10 and rule.str.startswith('C')")
    rows=[]
    for budget in [.01,.05,.10,.20]:
        row=['%d%%'%round(100*budget)]
        for pi in [.02,.05,.10,.15]:
            part=original_scen[(original_scen.budget==budget)&(original_scen.scenario_prevalence==pi)]
            check(len(part)==1,'unique budget/scenario source')
            r=part.iloc[0]
            # Reconstruct from empirical class-conditional rates rather than the derived difference column.
            loss=100*(10*pi*r.empirical_fnr+(1-pi)*r.empirical_fpr)-100*10*pi
            alerts=pi*(1-r.empirical_fnr)+(1-pi)*r.empirical_fpr
            true_mass=pi*(1-r.empirical_fnr)
            precision='NA' if alerts==0 else '%.2f%%'%(100*true_mass/alerts)
            row.append('%.2f\n%.2f%%\n%s'%(loss,100*alerts,precision))
        rows.append(row)
    check(getrows(16)==rows,'Table A5 paired values reconstructed independently')
    boundary_tables=[t for t in doc.tables if [c.text for c in t.rows[0].cells]==new_headers[1]]
    check(len(boundary_tables)==1,'one frozen-policy cost-boundary table')
    q=pooled.query("model=='XGBoost' and workflow=='fixed' and cost==10 and mechanism=='C'").sort_values('budget')
    rows=[]
    for _,r in q.iterrows():
        equality='—' if r.tp==0 else '%.2f'%(r.fp/r.tp)
        rows.append(['%d%%'%round(100*r.budget),str(int(r.tp)),str(int(r.fp)),str(int(r.alerts)),equality])
    check([[c.text for c in r.cells] for r in boundary_tables[0].rows][1:]==rows,
          'cost equality reconstructed from original pooled counts')
    links=[t for t in doc.tables if [c.text for c in t.rows[0].cells]==new_headers[0]]
    check(len(links)==1 and len(links[0].rows)==4,'one cross-stage evidence table')
    check('does not maintain a prescribed test alert rate' in md,'historical budget versus future fixed cutoff')
    check('historical selection uses rselect' in md.lower(),'selection cost distinct from evaluation cost')
    check('It is undefined when expected alert mass is zero' in md,'scenario precision zero-alert convention')
    check('Stage one [3] transfers a reference cutoff' in md,'cross-stage workflow link')
    check('budget cutoffs exclude boundary ties' in md,'cross-stage threshold distinction')
    check('Alert-count changes in the second panel are labelled as integers.' in md,'integer count caption')
    check('[4] Künsch HR.' in md,'block-bootstrap source')
    record=dict(passed=True,checks=checks,pages=len(pdf.pages),tables=len(doc.tables),figures=len(doc.inline_shapes),
                scope='Computed report tables, paragraphs, hashes, PDF decimals and package metadata; visual inspection is recorded separately.',

                report_sha256={p.name:sha(p) for p in dest.iterdir() if p.suffix in ('.docx','.pdf','.md')})
    return record


def verify_report(root, baseline, output):
    """Persist the current invocation before any input is read."""
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    started = datetime.now(timezone.utc).isoformat()
    invocation = dict(passed=False, status='running', started_at_utc=started,
                      verifier_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest())
    output.write_text(json.dumps(invocation, indent=2), encoding='utf-8')
    try:
        record = _run_checks(Path(root), Path(baseline))
    except (Exception, KeyboardInterrupt) as error:
        invocation.update(status='failed', failed_at_utc=datetime.now(timezone.utc).isoformat(),
                          error_type=type(error).__name__, error=str(error))
        output.write_text(json.dumps(invocation, indent=2), encoding='utf-8')
        raise
    record.update(status='passed', started_at_utc=started,
                  completed_at_utc=datetime.now(timezone.utc).isoformat(),
                  verifier_sha256=invocation['verifier_sha256'])
    output.write_text(json.dumps(record, indent=2), encoding='utf-8')
    print('Report checks passed:', record['checks'], 'pages', record['pages'])
    return record


def main():
    parser = argparse.ArgumentParser(description='Independent Stage 2 report-table checks')
    parser.add_argument('--root', type=Path, required=True)
    parser.add_argument('--baseline', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    verify_report(args.root, args.baseline, args.output)


if __name__ == '__main__':
    main()
